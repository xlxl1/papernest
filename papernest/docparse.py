"""多格式文档解析：Word / PowerPoint / HTML / Markdown / 纯文本 → 统一块结构 → Markdown。

为什么单独一个模块（而不是塞进 pdfimport）：
- pdfimport 那条链路（魔数校验 → 落盘 → 启发式抽元数据 → papers/pages 入库）是围绕 PDF
  定制的；这里只负责「文件 → 结构化文本」这一段，**不写盘、不写库、不联网、不调 LLM**。
  入库仍走主进程的既有路径（见模块末尾 `to_markdown` 的说明）。纯函数模块最好测，
  也最容易被别的场景（笔记导入、网页存档、组会 PPT 归档）复用。
- 批量导入 200 个文件时，真正稀缺的不是「解析得多漂亮」，而是**知道哪些没进来、为什么**。
  所以 `parse_batch` 的三态计数是本模块的主产出：
    skipped = 我们主动不收（格式不支持 / 空文件 / 超限 / 内容重复 / 抽不出任何文本）
    failed  = 该收但出错（reason 一定带上异常类型名，能直接拿去搜）
    ok      = 收下了
  把这两类混成一个 "error" 数字，等于把「你给的文件我不认」和「我崩了」揉在一起，
  排障时完全没法下手。

结构约定（所有抽取器统一返回）：
    {"text": str, "blocks": [block], "title": str|None, "meta": dict, "warnings": [str]}
    block = {"text": str, "kind": str, "index": int, ...}
    kind：heading（带 level）/ paragraph / table（带 rows）/ code（带 lang）
          / slide（带 title）/ notes
    index：除 pptx 外都是**块序号**（0 起）；**pptx 的 index 是页码**（1 起），
           同一页的 slide/table/notes 块共享同一个 index——这样任何一块都能定位到第几页，
           与本项目「证据带页码」的口径一致。

不编造原则在本模块的落点：抽不到标题就 title=None（绝不拿文件名或正文首行冒充），
编码探测走到兜底档就如实记 warning，解析出错就记异常类型而不是静默吞掉。
"""
import hashlib
import html as _html
import re
import zipfile
from html.parser import HTMLParser
from pathlib import Path

# ── 格式登记表 ──
# PDF 不在这里：主进程走既有 pdfimport 路径，本模块见到 PDF 会明确拒收并说明原因。
SUPPORTED = {
    ".docx": "word",
    ".pptx": "powerpoint",
    ".html": "html",
    ".htm": "html",
    ".md": "markdown",
    ".markdown": "markdown",
    ".txt": "text",
}

# 单文件体积上限，与 pdfimport.MAX_PDF_BYTES 对齐；超限记 skipped 而不是 failed
# （不是我们坏了，是我们主动不收）。
MAX_FILE_BYTES = 50 * 1024 * 1024

# 嗅探读取的文件头长度：4KB 足够覆盖 zip 魔数、PDF 魔数与 <!doctype html>。
_SNIFF_BYTES = 4096

_ZIP_MAGIC = b"PK\x03\x04"
_PDF_MAGIC = b"%PDF-"

# OOXML 各家的特征部件。xlsx 也列出来，是为了能明确回一句「这是 xlsx，不支持」，
# 而不是因为「也是个 zip」就误判成 word。
_ZIP_PARTS = (
    ("word/document.xml", "word"),
    ("ppt/presentation.xml", "powerpoint"),
    ("xl/workbook.xml", None),
)

# 编码探测链。utf-8-sig 必须排在 utf-8 前面：utf-8 也能解出带 BOM 的文件，
# 只是会在开头留一个 \ufeff，排在后面的 utf-8-sig 就永远轮不到了。
# GBK 排第三——中文环境下 GBK 文本文件极其常见，漏了它就是线上乱码。
TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gbk", "latin-1")

_WS_RE = re.compile(r"[ \t\r\f\v\u00a0\u3000]+")
_NL_RE = re.compile(r"\n{3,}")
# C0 \u63a7\u5236\u5b57\u7b26\uff08\u4fdd\u7559 \t \n \r\uff0c\u5b83\u4eec\u7531 _WS_RE/_NL_RE \u8d1f\u8d23\uff09\u3002NUL \u6df7\u8fdb\u6b63\u6587 \u2192 sqlite/FTS \u51fa\u4e8b\u3002
_CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0e-\x1f\x7f]")


def _collapse(s: str) -> str:
    """压缩行内空白但保留换行结构（表格单元格、段落内软换行都靠这个保形）。"""
    if not s:
        return ""
    s = _WS_RE.sub(" ", s.replace("\r\n", "\n").replace("\r", "\n"))
    s = "\n".join(line.strip() for line in s.split("\n"))
    return _NL_RE.sub("\n\n", s).strip()


def _blocks_text(blocks: list[dict]) -> str:
    return "\n\n".join(b["text"] for b in blocks if b.get("text"))


def _result(blocks: list[dict], title: str | None, meta: dict,
            warnings: list[str]) -> dict:
    return {"text": _blocks_text(blocks), "blocks": blocks,
            "title": title or None, "meta": meta, "warnings": warnings}


class _Emitter:
    """块累加器：统一分配 index，顺手挡掉空块（空块进了 RAG 就是纯噪声）。"""

    def __init__(self, start: int = 0, fixed_index: int | None = None):
        self.blocks: list[dict] = []
        self._n = start
        self._fixed = fixed_index

    def emit(self, kind: str, text: str, **extra) -> dict | None:
        # 控制字符统一在这里剔除（\n \t 不动）：二进制文件被改名成 .txt 时，NUL 会一路
        # 流进 blocks → sqlite → FTS。放在 emit 而不是 _collapse，是因为 code/pre 块
        # 走的是不经过 _collapse 的路径，那条路同样会漏。
        text = _CTRL_RE.sub("", (text or "")).strip("\n").rstrip()
        if not text.strip():
            return None
        idx = self._fixed if self._fixed is not None else self._n
        self._n += 1
        blk = {"text": text, "kind": kind, "index": idx}
        blk.update(extra)
        self.blocks.append(blk)
        return blk


# ── 编码探测 ──

def decode_bytes(data: bytes) -> tuple[str, str, list[str]]:
    """字节 → 文本，返回 (text, encoding, warnings)。

    latin-1 能解任何字节序列，所以它一旦命中就是「解得出，但多半是乱码」——
    必须如实记 warning，否则 GBK 文件被 latin-1 静默吃掉，入库的是一堆问号，
    而调用方以为一切正常。含 NUL 字节的（UTF-16/二进制）不给 latin-1 兜，
    直接走 replace 档并说明。
    """
    warnings: list[str] = []
    if not data:
        return "", "empty", warnings
    text: str | None = None
    encoding = ""
    # UTF-16 有 BOM 就直接认，不然 latin-1 会把它解成夹杂 \x00 的乱码
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        try:
            text, encoding = data.decode("utf-16"), "utf-16"
        except UnicodeDecodeError:
            pass
    if text is None:
        for enc in TEXT_ENCODINGS:
            if enc == "latin-1":
                if b"\x00" in data:
                    break  # 二进制/宽字符编码，latin-1 兜出来只会是垃圾
                warnings.append("编码探测未命中 utf-8/gbk，已按 latin-1 兜底解码，"
                                "非拉丁字符可能是乱码")
            try:
                text, encoding = data.decode(enc), enc
                break
            except UnicodeDecodeError:
                if enc == "latin-1":  # 理论上不会发生，留着防 codec 变更
                    warnings.pop()
                continue
    if text is None:
        warnings.append("编码探测全部失败（疑似二进制或未知编码），已用替换字符解码")
        text, encoding = data.decode("utf-8", errors="replace"), "utf-8/replace"
    # NUL 是合法的 UTF-8 码位，所以「含 NUL」并不会让上面的探测失败——二进制文件改名成
    # .txt 时会一路静默解码成功。必须在这里如实记账，否则调用方以为收到的是正常文本。
    if _CTRL_RE.search(text):
        warnings.append("解码结果含 NUL 等控制字符（多半是二进制文件改了扩展名），"
                        "这些字符已在分块时剔除")
    return text, encoding, warnings


def read_text(path) -> tuple[str, str, list[str]]:
    """读文件并探测编码。句柄用 with 关闭，本模块任何地方都不写文件。"""
    with open(path, "rb") as f:
        return decode_bytes(f.read())


# ── 格式识别：扩展名 + 魔数复核 ──

def _zip_kind(path) -> tuple[str | None, str]:
    """按 zip 内部部件判断是 docx 还是 pptx。返回 (format|None, 说明)。"""
    try:
        with zipfile.ZipFile(path) as z:
            names = set(z.namelist())
    except (zipfile.BadZipFile, OSError) as e:
        return None, f"ZIP 魔数对但打不开（{type(e).__name__}）"
    for part, fmt in _ZIP_PARTS:
        if part in names:
            if fmt:
                return fmt, f"ZIP 内含 {part}"
            return None, f"ZIP 内含 {part}（Excel 工作簿，本模块不支持）"
    return None, "ZIP 包但不含 Word/PowerPoint 部件"


def _looks_html(head: str, strict: bool) -> bool:
    """strict=True 只认不可能误判的强特征（给 .md 用，Markdown 里内联 HTML 是合法的）。"""
    low = head.lower()
    if any(m in low for m in ("<!doctype html", "<html", "<head", "<body")):
        return True
    return (not strict) and head.startswith("<")


def detect_format_detail(path) -> dict:
    """返回 {"format": str|None, "ext": str, "by": "ext"|"magic"|"none", "note": str}。

    改扩展名是常态（微信收到的 .docx 存成 .txt、导出的网页存成 .doc），所以
    **扩展名与内容矛盾时以内容为准**，并把这次改判写进 note，让调用方看得见。
    """
    p = Path(path)
    ext = p.suffix.lower()
    ext_fmt = SUPPORTED.get(ext)
    try:
        with open(p, "rb") as f:
            head = f.read(_SNIFF_BYTES)
    except OSError as e:
        return {"format": ext_fmt, "ext": ext, "by": "ext",
                "note": f"读文件头失败（{type(e).__name__}），只能按扩展名判断"}

    if not head:
        return {"format": ext_fmt, "ext": ext, "by": "ext", "note": "文件为空"}

    if head[:4] == _ZIP_MAGIC:
        zk, why = _zip_kind(p)
        if zk and zk != ext_fmt:
            return {"format": zk, "ext": ext, "by": "magic",
                    "note": f"扩展名 {ext or '(无)'} 与内容不符：{why}，以内容为准判为 {zk}"}
        if zk:
            return {"format": zk, "ext": ext, "by": "magic", "note": ""}
        return {"format": None, "ext": ext, "by": "magic", "note": why}

    if head[:len(_PDF_MAGIC)] == _PDF_MAGIC:
        return {"format": None, "ext": ext, "by": "magic",
                "note": "内容是 PDF，本模块不处理（请走 pdfimport 路径）"}

    text_head = head.decode("utf-8", errors="replace").lstrip("\ufeff \t\r\n")

    if ext_fmt in ("word", "powerpoint"):
        # 声称是 OOXML 却没有 zip 魔数：多半是别人把 txt/html 改名成了 .docx
        fallback = "html" if _looks_html(text_head, strict=False) else "text"
        return {"format": fallback, "ext": ext, "by": "magic",
                "note": f"扩展名 {ext} 声称是 {ext_fmt} 但内容不是 ZIP 包，"
                        f"以内容为准按 {fallback} 处理"}

    if ext_fmt == "html":
        if "<" not in text_head:
            return {"format": "text", "ext": ext, "by": "magic",
                    "note": f"扩展名 {ext} 但内容开头不含标签，以内容为准按纯文本处理"}
        return {"format": "html", "ext": ext, "by": "ext", "note": ""}

    if ext_fmt in ("markdown", "text"):
        if _looks_html(text_head, strict=(ext_fmt == "markdown")):
            return {"format": "html", "ext": ext, "by": "magic",
                    "note": f"扩展名 {ext} 但内容是 HTML，以内容为准按 html 处理"}
        return {"format": ext_fmt, "ext": ext, "by": "ext", "note": ""}

    # 扩展名不认识：只在内容有强特征时才收，否则不猜（.png/.exe 不该被当成文本读进来）
    if _looks_html(text_head, strict=True):
        return {"format": "html", "ext": ext, "by": "magic",
                "note": f"扩展名 {ext or '(无)'} 未登记，按内容判为 html"}
    return {"format": None, "ext": ext, "by": "none",
            "note": f"不支持的格式：{ext or '(无扩展名)'}"}


def detect_format(path) -> str | None:
    """先扩展名、再魔数复核；矛盾时以内容为准。矛盾说明见 detect_format_detail 的 note。"""
    return detect_format_detail(path)["format"]


def is_supported(path) -> bool:
    return detect_format(path) is not None


# ── Word ──

# 内置样式名在 styles.xml 里通常是英文，但中文版 Word 另存的文件里见过「标题 1」，
# 两种都认。Title/Subtitle 单独映射成 1/2 级。
_DOCX_HEADING_RE = re.compile(r"^(?:heading|标题)\s*([1-9])", re.I)


def _docx_heading_level(style_name: str | None) -> int | None:
    name = (style_name or "").strip()
    if not name:
        return None
    m = _DOCX_HEADING_RE.match(name)
    if m:
        return min(int(m.group(1)), 6)
    low = name.lower()
    if low == "title":
        return 1
    if low == "subtitle":
        return 2
    return None


def _iter_docx_body(doc, parent=None, depth: int = 0):
    """按文档 XML 顺序交替产出段落与表格。

    doc.paragraphs 与 doc.tables 是两个独立列表，先拼段落再拼表格会彻底丢掉
    「这张表夹在哪两段之间」——正文与表格错位，后面做证据定位就全是错的。

    还要下钻 `<w:sdt>`（内容控件）：Word 的自动目录、封面、模板占位区都把整段正文包在
    `<w:sdt><w:sdtContent>` 里。python-docx **生成**的文件永远不会有 sdt，所以只用
    程序化夹具测是看不出来的；真实投稿稿里丢掉的是整章内容，而且**一声不吭**。
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    el = doc.element.body if parent is None else parent
    for child in el.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield "p", Paragraph(child, doc)
        elif tag == "tbl":
            yield "tbl", Table(child, doc)
        elif tag in ("sdt", "sdtContent") and depth < 8:
            yield from _iter_docx_body(doc, child, depth + 1)


def _table_block_text(rows: list[list[str]]) -> str:
    """行内用 ` | ` 连接、行间换行——表头行留在第 0 行不动。

    表格被打散成一堆孤立单元格是 RAG 里最常见的质量杀手：脱离表头的 "0.83"
    没有任何检索价值。所以一张表恒定是**一个**块。
    """
    return "\n".join(" | ".join(c for c in row) for row in rows)


def _clean_rows(rows: list[list[str]]) -> list[list[str]]:
    return [r for r in rows if any(c.strip() for c in r)]


def extract_docx(path) -> dict:
    """python-docx 抽取：标题层级 + 表格整块 + core_properties。"""
    from docx import Document  # 局部导入：批量跑 .md 时不必付 docx 的导入成本
    warnings: list[str] = []
    doc = Document(str(path))
    em = _Emitter()
    first_heading = None

    for kind, obj in _iter_docx_body(doc):
        if kind == "p":
            text = _collapse(obj.text)
            if not text:
                continue
            try:
                level = _docx_heading_level(obj.style.name if obj.style else None)
            except (AttributeError, KeyError):  # 样式指向了不存在的 styleId
                level = None
            if level:
                blk = em.emit("heading", text, level=level)
                if blk and first_heading is None:
                    first_heading = text
            else:
                em.emit("paragraph", text)
        else:
            rows = []
            for row in obj.rows:
                try:
                    rows.append([_collapse(c.text) for c in row.cells])
                except (IndexError, ValueError, AttributeError) as e:
                    # 合并单元格结构异常时跳这一行，如实记账，别让整篇挂掉
                    warnings.append(f"表格某行读取失败（{type(e).__name__}），已跳过该行")
            rows = _clean_rows(rows)
            if rows:
                em.emit("table", _table_block_text(rows), rows=rows,
                        n_rows=len(rows), n_cols=max(len(r) for r in rows))

    meta = {"format": "word", "n_paragraphs": sum(
        1 for b in em.blocks if b["kind"] in ("paragraph", "heading"))}
    core_title = None
    try:
        cp = doc.core_properties
        for key in ("title", "author", "subject", "keywords", "last_modified_by"):
            val = getattr(cp, key, None)
            if val:
                meta[key] = str(val).strip()
        if cp.created:
            meta["created"] = cp.created.isoformat()
        core_title = (cp.title or "").strip() or None
    except (AttributeError, ValueError, KeyError) as e:
        warnings.append(f"core_properties 读取失败（{type(e).__name__}），元数据留空")

    # 标题只认「文档属性里的 title」或「第一个标题段」；正文首行不算——那是猜的
    return _result(em.blocks, core_title or first_heading, meta, warnings)


# ── PowerPoint ──

def _pptx_shape_items(shapes, depth: int = 0):
    """展开组合形状。深度设限，防畸形文件把递归打爆。"""
    for sh in shapes:
        if depth < 4 and getattr(sh, "shape_type", None) is not None \
                and hasattr(sh, "shapes"):
            yield from _pptx_shape_items(sh.shapes, depth + 1)
            continue
        yield sh


def extract_pptx(path) -> dict:
    """python-pptx 抽取：每页一个 slide 块（index=页码），备注单独成 notes 块。

    备注常常才是真正的解释性内容（正文只有几个词组），丢掉备注等于丢掉这页的语义。
    """
    from pptx import Presentation  # 局部导入，理由同 docx
    warnings: list[str] = []
    prs = Presentation(str(path))
    blocks: list[dict] = []
    first_title = None

    for page_no, slide in enumerate(prs.slides, 1):
        em = _Emitter(fixed_index=page_no)  # 同页各块共享 index=页码
        try:
            title_shape = slide.shapes.title
        except (AttributeError, ValueError):
            title_shape = None
        title_text = ""
        if title_shape is not None and title_shape.has_text_frame:
            title_text = _collapse(title_shape.text_frame.text)
        # python-pptx 每次访问都新建一个 proxy 对象，`shapes.title` 与遍历 `shapes` 拿到的
        # **不是同一个 Python 对象**（`is` 恒为 False），只有底层 XML 元素是同一个。
        # 用 `is` 比对象会导致标题被当成普通正文再收一遍，每页标题都重复一次。
        title_el = getattr(title_shape, "element", None) if title_shape is not None else None

        body_parts: list[str] = []
        tables: list[list[list[str]]] = []
        for sh in _pptx_shape_items(slide.shapes):
            if title_el is not None and getattr(sh, "element", None) is title_el:
                continue
            try:
                if getattr(sh, "has_table", False):
                    rows = _clean_rows([[_collapse(c.text) for c in row.cells]
                                        for row in sh.table.rows])
                    if rows:
                        tables.append(rows)
                    continue
                if getattr(sh, "has_text_frame", False):
                    t = _collapse(sh.text_frame.text)
                    if t:
                        body_parts.append(t)
            except (AttributeError, ValueError, KeyError) as e:
                warnings.append(f"第 {page_no} 页某个形状读取失败"
                                f"（{type(e).__name__}），已跳过")

        # slide 块是自包含的「一页」：标题行 + 正文；标题同时单独挂在 title 字段上
        page_text = "\n".join(x for x in ([title_text] + body_parts) if x)
        em.emit("slide", page_text, title=title_text or None,
                has_title=bool(title_text))
        for rows in tables:
            em.emit("table", _table_block_text(rows), rows=rows,
                    n_rows=len(rows), n_cols=max(len(r) for r in rows))
        try:
            if slide.has_notes_slide:
                notes = _collapse(slide.notes_slide.notes_text_frame.text)
                em.emit("notes", notes)
        except (AttributeError, ValueError) as e:
            warnings.append(f"第 {page_no} 页备注读取失败（{type(e).__name__}）")

        # 空白页不产生空块：em 内部已挡掉空文本，这里自然就是 0 块
        blocks.extend(em.blocks)
        if title_text and first_title is None:
            first_title = title_text

    meta = {"format": "powerpoint", "n_slides": len(prs.slides)}
    core_title = None
    try:
        cp = prs.core_properties
        for key in ("title", "author", "subject", "keywords"):
            val = getattr(cp, key, None)
            if val:
                meta[key] = str(val).strip()
        if cp.created:
            meta["created"] = cp.created.isoformat()
        core_title = (cp.title or "").strip() or None
    except (AttributeError, ValueError, KeyError) as e:
        warnings.append(f"core_properties 读取失败（{type(e).__name__}），元数据留空")

    return _result(blocks, core_title or first_title, meta, warnings)


# ── HTML ──

# 丢内容的标签，分两档：
# - 硬跳过：script/style 这类根本不是正文；它们的闭合标签在 html.parser 里是可靠的
#   （CDATA 模式一直读到 </script>），漏收的风险可以忽略。
# - 软跳过：nav/footer/aside 是导航样板，留着会把检索打偏（每篇都命中「首页 登录 关于我们」），
#   但它们是**普通元素**，手写页面里忘了闭合很常见。一旦忘闭合，整篇正文就被吞掉了——
#   所以见到 main/article/h1 这种「正文主体才会有」的标签，就判定软跳过区早该结束，
#   强制解除并记 warning。宁可混进一点样板，也不能整篇丢。
_HTML_HARD_SKIP = {"script", "style", "noscript", "template", "svg", "iframe"}
_HTML_SOFT_SKIP = {"nav", "footer", "aside"}
_HTML_SKIP_TAGS = _HTML_HARD_SKIP | _HTML_SOFT_SKIP
_HTML_SOFT_SKIP_BREAKERS = {"main", "article", "h1"}
_HTML_BLOCK_TAGS = {"p", "div", "section", "article", "header", "main", "li",
                    "ul", "ol", "dl", "dt", "dd", "blockquote", "hr",
                    "figure", "figcaption", "address", "form", "tr"}
_HTML_HEADING_RE = re.compile(r"^h([1-6])$")


class _HtmlExtractor(HTMLParser):
    """标准库 html.parser（bs4 没装）。

    两条硬要求：
    1) **标签不闭合不能崩**——网页存档里未闭合的 <p>/<td>/<li> 是常态。所以这里不维护
       通用标签栈做配对，只维护「跳过区」「表格」「标题」三个局部状态，任何一个都
       只在能识别的边界上开合，识别不到就当没发生。
    2) **实体必须反转义**：convert_charrefs=False + handle_entityref/charref 显式
       html.unescape 一次。用默认的 True 再 unescape 一遍，会把 `&amp;lt;` 这种
       转义过两轮的文本错误地还原成 `<`。
    """

    def __init__(self):
        super().__init__(convert_charrefs=False)
        self.em = _Emitter()
        self.title: str | None = None
        self.meta: dict = {"format": "html"}
        self.warnings: list[str] = []
        self._buf: list[str] = []
        self._skip: list[str] = []
        self._heading: tuple[int, list[str]] | None = None
        self._tables: list[dict] = []
        self._title_buf: list[str] | None = None
        self._pre = 0

    # -- 文本落点 --
    def _sink(self, s: str):
        if self._skip or not s:
            return
        if self._title_buf is not None:
            self._title_buf.append(s)
            return
        t = self._tables[-1] if self._tables else None
        if t is not None:
            if t["cell"] is not None:
                t["cell"].append(s)
            return  # 表格内、单元格外的游离文本（多半是空白）直接丢
        if self._heading is not None:
            self._heading[1].append(s)
            return
        self._buf.append(s)

    def handle_data(self, data):
        self._sink(data)

    def handle_entityref(self, name):
        self._sink(_html.unescape(f"&{name};"))

    def handle_charref(self, name):
        self._sink(_html.unescape(f"&#{name};"))

    # -- 段落/标题/表格 --
    def _commit_title(self, forced: bool = False):
        """把 <title> 缓冲区落地。

        HTML5 里 <title> 只允许纯文本，所以在 title 里见到**任何**标签都说明 </title>
        丢了。不强制收口的话，`_sink` 会把整篇正文都倒进 title 缓冲区，最终产出
        0 个块——`parse()` 只会报一句 skipped「内容为空」，整份网页存档就这么无声消失了。
        这和 <nav> 忘闭合吞掉正文是同一类事故，处理方式也保持一致：强制收口 + 记 warning。
        """
        if self._title_buf is None:
            return
        self.title = _collapse("".join(self._title_buf)) or None
        self._title_buf = None
        if forced:
            self.warnings.append("未闭合的 <title>，已在下一个标签处强制收口")

    def _close_heading(self):
        """标题在任何块级边界处收尾。

        `<h1>标题<p>正文` 这种漏了 </h1> 的写法很常见；不在块级边界收口的话，
        整篇正文会被并进标题块里（实测踩到过）。
        """
        if self._heading is None:
            return
        level, frags = self._heading
        self._heading = None
        self.em.emit("heading", _collapse("".join(frags)), level=level)

    def _flush_para(self):
        """<pre> 里的内容按 code 块原样保留：把缩进和换行压掉，代码就没法看了。"""
        self._close_heading()
        raw = "".join(self._buf)
        self._buf = []
        if self._pre:
            self.em.emit("code", raw.strip("\n"), lang=None)
        else:
            self.em.emit("paragraph", _collapse(raw))

    def _close_cell(self):
        t = self._tables[-1] if self._tables else None
        if t is None or t["cell"] is None:
            return
        if not t["rows"]:
            t["rows"].append([])
        t["rows"][-1].append(_collapse("".join(t["cell"])))
        t["cell"] = None

    def _close_table(self):
        if not self._tables:
            return
        self._close_cell()
        t = self._tables.pop()
        rows = _clean_rows(t["rows"])
        if rows:
            self.em.emit("table", _table_block_text(rows), rows=rows,
                         n_rows=len(rows), n_cols=max(len(r) for r in rows))

    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        # 放在最前面：<title> 未闭合时，后面无论出现什么标签都必须先把 title 收口，
        # 否则连 <body>/<script> 都会被当成标题文字继续往缓冲区里灌。
        if self._title_buf is not None and tag != "title":
            self._commit_title(forced=True)
        if tag in _HTML_SKIP_TAGS:
            self._skip.append(tag)
            return
        if self._skip:
            if tag in _HTML_SOFT_SKIP_BREAKERS \
                    and all(t in _HTML_SOFT_SKIP for t in self._skip):
                self.warnings.append(
                    f"未闭合的 <{'>/<'.join(self._skip)}>，已在 <{tag}> 处强制结束跳过区")
                self._skip = []
            else:
                return
        if tag == "title":
            self._title_buf = []
            return
        if tag == "meta":
            d = {(k or "").lower(): (v or "") for k, v in attrs}
            name = (d.get("name") or d.get("property") or "").lower()
            if name in ("description", "author", "og:title") and d.get("content"):
                self.meta.setdefault(name.replace("og:", ""), d["content"].strip())
            return
        if tag == "br":
            self._sink("\n")
            return
        if tag == "table":
            self._flush_para()
            self._tables.append({"rows": [], "cell": None})
            return
        if self._tables:
            if tag == "tr":
                self._close_cell()
                self._tables[-1]["rows"].append([])
                return
            if tag in ("td", "th"):
                self._close_cell()
                if not self._tables[-1]["rows"]:
                    self._tables[-1]["rows"].append([])
                self._tables[-1]["cell"] = []
                return
            return  # 表格内的其它标签不参与分块，文本继续进当前单元格
        m = _HTML_HEADING_RE.match(tag)
        if m:
            self._flush_para()
            self._heading = (int(m.group(1)), [])
            return
        if tag == "pre":
            self._flush_para()
            self._pre += 1
            return
        if tag in _HTML_BLOCK_TAGS:
            self._flush_para()

    def handle_endtag(self, tag):
        tag = tag.lower()
        # `<title>abc</head>` 这类只有闭标签、没有后续开标签的写法同样要能收口
        if self._title_buf is not None and tag != "title":
            self._commit_title(forced=True)
        if tag in _HTML_SKIP_TAGS:
            if tag in self._skip:  # 就近弹栈，配不上就当没看见（不闭合是常态）
                del self._skip[len(self._skip) - 1 - self._skip[::-1].index(tag):]
            return
        if tag in ("body", "html") and self._skip:
            # 安全阀：<nav> 之类忘了闭合会把后面整篇吞掉，见到 </body> 强制解除跳过
            self.warnings.append(f"存在未闭合的 {'/'.join(self._skip)} 标签，"
                                 f"已在 </{tag}> 处强制结束跳过区")
            self._skip = []
        if self._skip:
            return
        if tag == "title":
            self._commit_title()
            return
        if tag == "table":
            self._close_table()
            return
        if self._tables:
            if tag in ("td", "th"):
                self._close_cell()
            return
        if _HTML_HEADING_RE.match(tag) and self._heading is not None:
            self._close_heading()
            return
        if tag == "pre":
            self._flush_para()
            self._pre = max(0, self._pre - 1)
            return
        if tag in _HTML_BLOCK_TAGS:
            self._flush_para()

    def finish(self):
        """收尾：未闭合的表格/标题/段落都要落地，不能因为缺个 </table> 就丢内容。"""
        self._commit_title(forced=True)   # `<title>abc` 直接到 EOF 的截断存档
        while self._tables:
            self.warnings.append("存在未闭合的 <table>，已按已读到的行输出")
            self._close_table()
        self._flush_para()   # 内部会先给未闭合的标题收尾


_TITLE_OPEN_RE = re.compile(r"<title\b[^>]*>", re.I)
_TITLE_CLOSE_RE = re.compile(r"</title", re.I)


def _repair_unclosed_title(text: str) -> tuple[str, bool]:
    """`<title>` 忘了闭合时，在下一个标签前补上 `</title>`。

    这一步必须在喂给 html.parser **之前**做，不能只靠解析回调兜：
    Python 3.13 起 html.parser 把 `<title>` 当 RCDATA（`RCDATA_CONTENT_ELEMENTS`），
    一旦缺 `</title>`，剩下整篇文档都会被当成一段纯文本吐回来，**后面一个 starttag
    回调都不会触发**——解析器层面的收口逻辑根本没有机会执行。
    Python 3.12（线上 Docker 镜像的版本）没有这条 RCDATA 规则，走的是另一条路径，
    所以两层防护都要留着，不然换个 Python 版本行为就不一样。
    """
    m = _TITLE_OPEN_RE.search(text)
    if not m or _TITLE_CLOSE_RE.search(text, m.end()):
        return text, False
    nxt = text.find("<", m.end())
    if nxt == -1:
        return text + "</title>", True
    return text[:nxt] + "</title>" + text[nxt:], True


def extract_html(path_or_text) -> dict:
    """接受**文件路径或 HTML 文本**。含 `<` 的一律当文本（路径里不会有尖括号）。"""
    src = str(path_or_text)
    warnings: list[str] = []
    meta_extra: dict = {}
    if "<" in src:
        text = src
    else:
        text, enc, warnings = read_text(src)
        meta_extra["encoding"] = enc

    text, repaired = _repair_unclosed_title(text)
    if repaired:
        warnings = warnings + ["未闭合的 <title>，已在下一个标签前补上 </title>"
                               "（否则整篇正文会被当成标题吞掉）"]

    p = _HtmlExtractor()
    try:
        p.feed(text)
        p.close()
    except Exception as e:  # html.parser 在畸形输入上不该抛，抛了也只降级不上升
        p.warnings.append(f"HTML 解析中断（{type(e).__name__}: {e}），已输出已解析部分")
    p.finish()
    p.meta.update(meta_extra)
    p.meta["n_blocks"] = len(p.em.blocks)
    return _result(p.em.blocks, p.title, p.meta, warnings + p.warnings)


# ── Markdown ──

_MD_FENCE_RE = re.compile(r"^\s{0,3}(`{3,}|~{3,})\s*([^`]*)$")
_MD_ATX_RE = re.compile(r"^\s{0,3}(#{1,6})(?:\s+(.*?))?\s*#*\s*$")
_MD_SETEXT_H1_RE = re.compile(r"^\s{0,3}=+\s*$")
_MD_HR_RE = re.compile(r"^\s{0,3}((-\s*){3,}|(\*\s*){3,}|(_\s*){3,})$")
# 分隔行 |---|:--:|。单列表格（|---|）也要认，所以第一段之后的列是可选重复。
# 它能匹配裸的 "---"，但调用点只在「本行含 |」时才问它，所以分隔线不会被误吃。
_MD_TABLE_SEP_RE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


# 只在**未转义**的 | 处切列。to_markdown 会把单元格里的 | 写成 \|，如果这里照旧
# 无脑 split("|")，`to_markdown → extract_markdown` 这条回路就会把一个两列的表切成
# 三列（["a|b","c"] → ["a\\","b","c"]）。to_markdown 的产物是本模块对下游承诺的
# 「统一中间形态」，它必须能被自己再读回来。
_MD_CELL_SPLIT_RE = re.compile(r"(?<!\\)\|")


def _md_row(line: str) -> list[str]:
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    return [c.strip().replace("\\|", "|") for c in _MD_CELL_SPLIT_RE.split(s)]


def _md_is_table_line(line: str) -> bool:
    return line.strip().startswith("|")


def _parse_frontmatter(lines: list[str]) -> tuple[dict, int]:
    """YAML frontmatter 的极简读取（Obsidian 笔记里到处都是）。

    只认 `key: value`，不引 YAML 解析器；100 行内找不到闭合分隔符就判定不是
    frontmatter（避免把正文里的 --- 当成开头而吃掉半篇文章）。
    """
    if not lines or lines[0].strip() not in ("---", "+++"):
        return {}, 0
    delim = lines[0].strip()
    for i in range(1, min(len(lines), 101)):
        if lines[i].strip() == delim:
            fm: dict = {}
            for raw in lines[1:i]:
                if ":" in raw and not raw.strip().startswith("#"):
                    k, v = raw.split(":", 1)
                    k, v = k.strip(), v.strip().strip("\"'")
                    if k:
                        fm[k] = v
            return fm, i + 1
    return {}, 0


def extract_markdown(path) -> dict:
    """`#` 层级 → heading；``` 围栏代码整块保留；`|` 行 → table。

    代码块必须整块留着：把 ```python 里的 `# comment` 当成一级标题、
    再把代码按空行切成碎段，是 Markdown 解析最典型的自伤。
    """
    src = str(path)
    # 「多行 + 不是一个真实存在的文件」才当成直接喂进来的文本。
    # 原判据是「多行 + Path(src).suffix 为空」，可 Path("# 标题\n\n详见 report.md\n").suffix
    # 是 ".md\n"——正文最后一行只要长得像文件名，就会被拿去 open() 并抛 OSError。
    # Path.is_file() 对非法路径只会返回 False，不会抛。
    if "\n" in src and not Path(src).is_file():
        text, enc, warnings = src, "inline", []  # 允许直接喂文本，便于测试与复用
    else:
        text, enc, warnings = read_text(src)
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")

    fm, start = _parse_frontmatter(lines)
    em = _Emitter()
    meta: dict = {"format": "markdown", "encoding": enc}
    if fm:
        meta["frontmatter"] = fm
    title = (fm.get("title") or "").strip() or None
    first_h1 = None

    buf: list[str] = []

    def flush():
        nonlocal buf
        if buf:
            em.emit("paragraph", _collapse("\n".join(buf)))
            buf = []

    i = start
    n = len(lines)
    while i < n:
        line = lines[i]

        m = _MD_FENCE_RE.match(line)
        if m:
            fence, info = m.group(1), (m.group(2) or "").strip()
            body: list[str] = []
            i += 1
            closed = False
            while i < n:
                m2 = _MD_FENCE_RE.match(lines[i])
                if m2 and m2.group(1)[0] == fence[0] and len(m2.group(1)) >= len(fence) \
                        and not (m2.group(2) or "").strip():
                    closed = True
                    i += 1
                    break
                body.append(lines[i])
                i += 1
            if not closed:
                warnings.append("代码块围栏未闭合，已把剩余内容整体作为代码块")
            flush()
            em.emit("code", "\n".join(body), lang=info or None)
            continue

        m = _MD_ATX_RE.match(line)
        if m:
            flush()
            level = len(m.group(1))
            blk = em.emit("heading", _collapse(m.group(2) or ""), level=level)
            if blk and level == 1 and first_h1 is None:
                first_h1 = blk["text"]
            i += 1
            continue

        # setext 一级标题（=== 下划线）。--- 不认：它同时是分隔线与表格分隔行，歧义太大
        if buf and len(buf) == 1 and _MD_SETEXT_H1_RE.match(line) and buf[0].strip():
            text_h = _collapse(buf[0])
            buf = []
            blk = em.emit("heading", text_h, level=1)
            if blk and first_h1 is None:
                first_h1 = blk["text"]
            i += 1
            continue

        if _MD_HR_RE.match(line):
            flush()
            i += 1
            continue

        nxt = lines[i + 1] if i + 1 < n else ""
        starts_table = _md_is_table_line(line) or (
            "|" in line and "|" in nxt and _MD_TABLE_SEP_RE.match(nxt))
        if starts_table:
            flush()
            rows: list[list[str]] = []
            while i < n and lines[i].strip() and "|" in lines[i]:
                if not _MD_TABLE_SEP_RE.match(lines[i]):
                    rows.append(_md_row(lines[i]))
                i += 1
            rows = _clean_rows(rows)
            if rows:
                em.emit("table", _table_block_text(rows), rows=rows,
                        n_rows=len(rows), n_cols=max(len(r) for r in rows))
            continue

        if not line.strip():
            flush()
            i += 1
            continue

        buf.append(line)
        i += 1

    flush()
    return _result(em.blocks, title or first_h1, meta, warnings)


# ── 纯文本 ──

def extract_text(path) -> dict:
    """按空行分段。编码探测见 decode_bytes——GBK 是中文环境最常见的线上翻车点。"""
    text, enc, warnings = read_text(path)
    em = _Emitter()
    for para in re.split(r"\n[ \t]*\n+", text.replace("\r\n", "\n").replace("\r", "\n")):
        em.emit("paragraph", _collapse(para))
    meta = {"format": "text", "encoding": enc,
            "n_lines": text.count("\n") + (1 if text else 0)}
    if em.blocks:
        # 首行留给调用方参考，但**不**当标题——那是猜的
        meta["first_line"] = em.blocks[0]["text"].split("\n")[0][:200]
    return _result(em.blocks, None, meta, warnings)


_EXTRACTORS = {
    "word": extract_docx,
    "powerpoint": extract_pptx,
    "html": extract_html,
    "markdown": extract_markdown,
    "text": extract_text,
}


# ── 单文件顶层入口 ──

def _empty(path, fmt, status, reason, warnings=None) -> dict:
    return {"path": str(path), "format": fmt, "status": status, "reason": reason,
            "title": None, "text": "", "blocks": [], "meta": {},
            "warnings": warnings or [], "n_blocks": 0, "n_chars": 0}


def parse(path) -> dict:
    """单文件解析：detect → 分派 → 统一结构。**任何情况都不抛异常**。

    返回值在成功与失败时键完全一致（调用方不必写两套取值逻辑）：
      path / format / status / reason / title / text / blocks / meta / warnings
      / n_blocks / n_chars
    status 三态：ok（收下了）、skipped（主动不收）、failed（该收但出错，reason 带异常类型）。
    """
    p = Path(path)
    try:
        if not p.exists():
            return _empty(p, None, "failed", "文件不存在")
        if not p.is_file():
            return _empty(p, None, "failed", "不是普通文件（目录或设备文件）")
        size = p.stat().st_size
    except OSError as e:
        return _empty(p, None, "failed", f"无法访问文件（{type(e).__name__}: {e}）")

    if size == 0:
        return _empty(p, SUPPORTED.get(p.suffix.lower()), "skipped", "空文件")
    if size > MAX_FILE_BYTES:
        return _empty(p, SUPPORTED.get(p.suffix.lower()), "skipped",
                      f"文件 {size / 1048576:.1f}MB 超过 "
                      f"{MAX_FILE_BYTES // 1048576}MB 上限")

    detail = detect_format_detail(p)
    fmt, note = detail["format"], detail["note"]
    warnings = [note] if note else []
    if not fmt:
        # 「扩展名本该支持、但内容验不过」是**失败**，不是「格式不支持」：
        # 一份损坏的 .docx 记成 skipped 会让批量报告把真正的坏文件混进
        # 「我主动没收」那一堆里，静默丢数据。只有扩展名本身就不支持才是 skipped。
        claimed = SUPPORTED.get(p.suffix.lower())
        status = "failed" if claimed else "skipped"
        return _empty(p, claimed, status, note or "不支持的格式", warnings)

    try:
        out = _EXTRACTORS[fmt](str(p))
    except Exception as e:
        # 解析器再健壮也会有意外（加密的 docx、损坏的 zip）。reason 必须带异常类型，
        # 否则 200 个文件的批量报告里只会看到一片「解析失败」，没法归因。
        return _empty(p, fmt, "failed", f"{type(e).__name__}: {e}", warnings)

    blocks = out.get("blocks") or []
    text = out.get("text") or ""
    res = {"path": str(p), "format": fmt, "status": "ok", "reason": "",
           "title": out.get("title"), "text": text, "blocks": blocks,
           "meta": out.get("meta") or {},
           "warnings": warnings + list(out.get("warnings") or []),
           "n_blocks": len(blocks), "n_chars": len(text)}
    if not blocks or not text.strip():
        res["status"] = "skipped"
        res["reason"] = "内容为空（未抽到任何文本块）"
    return res


# ── 批量：三态计数 ──

def _sha256_file(path) -> str | None:
    h = hashlib.sha256()
    try:
        # 去重哈希跑在 parse() 的体积闸门**之前**，不设上限的话，一个被误拖进目录的
        # 4GB 视频会被完整读一遍，只为了拿一个马上就要丢掉的哈希。超限文件不参与去重，
        # 反正 parse() 接下来就会把它记成 skipped(超限)。
        if Path(path).stat().st_size > MAX_FILE_BYTES:
            return None
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(1 << 20), b""):
                h.update(chunk)
    except OSError:
        return None
    return h.hexdigest()


def _bump(by_format: dict, fmt: str, status: str):
    slot = by_format.setdefault(fmt, {"ok": 0, "failed": 0, "skipped": 0})
    slot[status] = slot.get(status, 0) + 1


def parse_batch(paths, progress=None, keep_parsed: bool = False) -> dict:
    """批量解析，返回三态计数 + 逐文件明细 + 按格式分组。

    - **单个文件失败绝不中断整批**：所有异常在 parse() 里就被收成 status=failed。
    - 内容重复（sha256 相同）判 skipped，reason 指出与哪个文件重复——同一份文件
      换个名字混在批次里是常态，让它进两遍等于污染库。
    - progress(frac, stage, message) 可选；回调自己炸了也不许影响解析
      （进度条不是业务）。
    - keep_parsed=True 时每个 item additional 带上完整 parsed 结果，供调用方直接入库；
      默认 False，避免 200 个文件的全文同时驻留内存。
    """
    items: list[dict] = []
    counts = {"ok": 0, "failed": 0, "skipped": 0}
    by_format: dict = {}
    seen: dict[str, str] = {}
    paths = list(paths or [])
    total = len(paths)

    def _tick(frac: float, stage: str, message: str):
        if not progress:
            return
        try:
            progress(max(0.0, min(1.0, frac)), stage, message)
        except Exception:
            pass  # 回调是观测手段，不能反过来搞挂解析

    for i, raw in enumerate(paths):
        name = Path(str(raw)).name or str(raw)
        _tick(i / total if total else 1.0, "parse", f"解析 {name}（{i + 1}/{total}）")

        sha = _sha256_file(raw)
        if sha and sha in seen:
            res = _empty(raw, detect_format_detail(raw)["format"], "skipped",
                         f"重复内容（与 {seen[sha]} 的 sha256 相同）")
        else:
            res = parse(raw)
            if sha and res["status"] == "ok":
                seen[sha] = str(Path(str(raw)).name)

        fmt = res["format"] or "unknown"
        counts[res["status"]] += 1
        _bump(by_format, fmt, res["status"])
        item = {"path": str(raw), "status": res["status"], "format": fmt,
                "reason": res["reason"], "n_blocks": res["n_blocks"],
                "n_chars": res["n_chars"], "title": res["title"],
                "warnings": res["warnings"]}
        if keep_parsed:
            item["parsed"] = res
        items.append(item)

    _tick(1.0, "done",
          f"完成：成功 {counts['ok']} / 失败 {counts['failed']} / 跳过 {counts['skipped']}")
    return {"ok": counts["ok"], "failed": counts["failed"],
            "skipped": counts["skipped"], "items": items, "by_format": by_format}


# ── 统一中间形态 ──

def _md_escape_cell(s: str) -> str:
    return s.replace("|", "\\|").replace("\n", "<br>")


def _md_table(rows: list[list[str]]) -> str:
    """行列补齐再输出，缺列的表格在 Markdown 里会渲染错位。"""
    width = max(len(r) for r in rows)
    out = []
    for j, row in enumerate(rows):
        cells = [_md_escape_cell(c) for c in row] + [""] * (width - len(row))
        out.append("| " + " | ".join(cells) + " |")
        if j == 0:  # 保留表头行，分隔行紧随其后
            out.append("| " + " | ".join(["---"] * width) + " |")
    return "\n".join(out)


def to_markdown(parsed: dict) -> str:
    """blocks → Markdown。这是喂给下游（切块 / 入库 / 卡片生成）的统一中间形态。"""
    if not parsed:
        return ""
    blocks = parsed.get("blocks") or []
    parts: list[str] = []
    title = (parsed.get("title") or "").strip()
    first = blocks[0] if blocks else None
    if title and not (first and first.get("kind") == "heading"
                      and first.get("text", "").strip() == title):
        parts.append(f"# {title}")

    for b in blocks:
        kind, text = b.get("kind"), (b.get("text") or "").strip()
        if not text:
            continue
        if kind == "heading":
            level = max(1, min(int(b.get("level") or 1), 6))
            parts.append("#" * level + " " + text.replace("\n", " "))
        elif kind == "table":
            rows = b.get("rows")
            parts.append(_md_table(rows) if rows else text)
        elif kind == "slide":
            head = f"## 第 {b.get('index')} 页"
            stitle = (b.get("title") or "").strip()
            if stitle:
                head += f"：{stitle}"
                # 标题已在小节名里，正文里就别重复一遍。**只去掉首行**：page_text 就是
                # 「标题行 + 正文」拼出来的，标题只可能在首行。原来是过滤掉所有等于标题
                # 的行，那会顺手删掉正文里合法重复的一行（比如「结论」页正文里也写了
                # 「结论」），而且正好把 extract_pptx 的标题重复 bug 掩盖成看不出来。
                lines = text.split("\n")
                body = "\n".join(lines[1:] if lines[0].strip() == stitle else lines)
            else:
                body = text
            parts.append(head)
            if body.strip():
                parts.append(body)
        elif kind == "notes":
            parts.append("\n".join(f"> 备注：{ln}" if k == 0 else f"> {ln}"
                                   for k, ln in enumerate(text.split("\n"))))
        elif kind == "code":
            lang = b.get("lang") or ""
            parts.append(f"```{lang}\n{text}\n```")
        else:
            parts.append(text)
    return "\n\n".join(parts).strip() + ("\n" if parts else "")
