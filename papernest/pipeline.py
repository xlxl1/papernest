"""多 Agent 写作流水线：选题 → 大纲 → 撰写⇄文献 → 润色 → 机械终检。

五个 Agent 是职责分离（各自 prompt / 产物落库 / 缓存键 / 模型档位），不是自由对话：
- 编排是确定性固定 DAG + 两个人工检查点（定题、改纲），复用既有任务基建
  （每个阶段 = jobs.kind='write' 的一个异步任务；检查点 = 任务结束、等用户指令）；
- 文献 Agent 是 cite.recommend 在写作循环内的化身——引用白名单制：
  生成稿每条 [n] 必须来自「证据句机械回取校验通过」的文献，防引用幻觉下移到写作；
- 节级内容寻址缓存：hash(prompt 版本, 标题, 论点, 术语表) 为键（不含大纲版本——
  改别处不影响本节），重跑未变节 0 token；单节强制重跑绕过缓存。
- 学术诚信红线：不生成实验数据；无证据处标「待补证据」；引用只来自真实库内文献。
"""
from __future__ import annotations

import hashlib
import json
import re
import uuid

from . import cards, cite, config, db, embeddings, llm, writing

PROMPT_VER = "p2v1"          # 写作 prompt 版本：升级 prompt 后缓存自然失效
CHECKPOINT_THRESHOLD = 7     # 节评审分 < 7 触发重写
TOPIC_SYSTEM = """你是科研选题顾问。基于给定文献卡片（tldr / 局限 / 摘要）为用户课题提出候选研究方向。
要求：
- 每个候选是一个具体、可写成综述的研究缺口，指名哪些文献的局限支撑它；
- evidence 里的 quote 必须逐字摘自对应文献卡片的「局限」或「摘要」原文，一字不改；
- 输出 JSON：{"topics":[{"title":"题目","gap":"缺口陈述","angle":"切入角度",
  "risks":"主要风险","evidence":[{"paper_id":123,"quote":"原文句子"}]}]}，共 3-4 个候选。
只输出 JSON。"""

OUTLINE_SYSTEM = """你是论文大纲规划师。基于选题与给定文献卡片产出结构化大纲 JSON：
{"title":"全文标题","sections":[{"no":1,"title":"节标题","points":["论点句",…],
"paper_ids":[12,15],"words":600}]}
约束：
- 4-6 节，各节 words 之和接近目标字数；
- paper_ids 只能从给定文献里选，points 里写清用它们论证什么；
- 确实没有文献支撑的节写 "no_support": true（诚实优于编造）。
只输出 JSON。"""

WRITER_SYSTEM = """你是严谨的学术写作者，负责撰写论文的一个章节。
铁律：
- 引用白名单是唯一引用来源，格式 [n]，n 是白名单序号；白名单之外的编号一律不许出现；
- 白名单文献支撑不了的说法不要写成事实：要么删掉，要么句尾标「（待补证据）」；
- 严禁编造数字、实验结果或白名单之外的文献；
- 用词与给定术语表一致，与其它章节自然衔接；
- 输出本节 Markdown 正文（不要输出章节标题，不要重复任务描述）。"""


# ── 运行记录 CRUD ──

def create_run(topic: str, project_id: int | None = None,
               target_words: int = 3000, max_rewrites: int = 2) -> dict:
    topic = (topic or "").strip()
    if len(topic) < 4:
        raise ValueError("课题描述太短（至少 4 个字）")
    run_id = uuid.uuid4().hex[:12]
    db.init_db()
    with db.conn() as c:
        c.execute("""INSERT INTO writing_runs(id,project_id,topic,config_json)
                     VALUES(?,?,?,?)""",
                  (run_id, project_id, topic,
                   json.dumps({"target_words": target_words,
                               "max_rewrites": max(0, min(int(max_rewrites), 3))},
                              ensure_ascii=False)))
        c.commit()
    return get_run(run_id)


def _row_to_run(r) -> dict:
    d = dict(r)
    d["topics"] = json.loads(d.pop("topics_json") or "[]")
    d["glossary"] = json.loads(d.pop("glossary_json") or "{}")
    d["config"] = json.loads(d.pop("config_json") or "{}")
    if d.get("result_json"):
        d["result"] = json.loads(d.pop("result_json"))
    else:
        d.pop("result_json")
    return d


def get_run(run_id: str) -> dict | None:
    with db.conn() as c:
        r = c.execute("SELECT * FROM writing_runs WHERE id=?", (run_id,)).fetchone()
    return _row_to_run(r) if r else None


def list_runs(limit: int = 20) -> list[dict]:
    db.init_db()
    with db.conn() as c:
        rows = c.execute("""SELECT id, topic, title, status, stage, created_at, updated_at
                            FROM writing_runs ORDER BY created_at DESC LIMIT ?""",
                         (limit,)).fetchall()
    return [dict(r) for r in rows]


_COLS = {"topics": "topics_json", "glossary": "glossary_json",
         "config": "config_json", "result": "result_json"}


def _update_run(run_id: str, **fields):
    sets, vals = [], []
    for k, v in fields.items():
        sets.append(f"{_COLS.get(k, k)}=?")
        vals.append(json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v)
    sets.append("updated_at=datetime('now','localtime')")
    with db.conn() as c:
        c.execute(f"UPDATE writing_runs SET {', '.join(sets)} WHERE id=?", (*vals, run_id))
        c.commit()


# ── Agent 0：文献 Agent（引用白名单）─────────────────────────────

def _verify_quote(paper_id, quote) -> bool:
    """机械回取：quote 必须逐字存在于该文献的卡片/摘要原文（归一化后包含）。

    两个入参都来自**模型输出**，类型不可信：实测模型给过 `quote: 123`
    （AttributeError: 'int' has no 'strip'）和 `paper_id: "abc"`（int() ValueError）。
    这里就地收口，不让类型问题冒到调用方去。
    """
    if not isinstance(quote, str) or not quote.strip():
        return False
    try:
        paper_id = int(paper_id)
    except (TypeError, ValueError):
        return False
    with db.conn() as c:
        r = c.execute("SELECT abstract, card_json FROM papers WHERE id=?",
                      (paper_id,)).fetchone()
    if not r:
        return False
    try:
        card = json.loads(r["card_json"] or "{}")
    except json.JSONDecodeError:
        card = {}
    haystack = " ".join(str(card.get(k) or "") for k in
                        ("tldr", "problem", "method", "results", "limitations"))
    haystack += " " + (r["abstract"] or "")
    return cite._norm(quote) in cite._norm(haystack)


def _whitelist(text: str, k: int = 6,
               extra_ids: list[int] | None = None) -> tuple[list[dict], str | None]:
    """文献 Agent：段落 → 引用白名单。返回 (白名单, 构造失败原因)。

    只收「证据句通过机械回取校验」的文献（cite.recommend 的 verified 口径）；
    大纲指定但推荐未命中的文献单独核验后补入。白名单即撰写 Agent 的引用边界。

    失败原因必须单独返回：「cite.recommend 抛异常」和「库里真的没有可支撑文献」
    原来产出的东西一模一样，都渲染成「本节无支撑」写进论文草稿并落 sections 表。
    前者是需要修的故障，后者是应当如实交付的事实——一次向量服务抖动就会让某一节
    永久变成无引用段，而没有任何地方记下这件事发生过。
    """
    text = (text or "").strip()
    if not text:
        return [], None
    out: dict[int, dict] = {}
    error: str | None = None
    try:
        for cand in cite.recommend(text, k * 2)["candidates"]:
            if cand["verified"] and cand["paper_id"] not in out:
                out[cand["paper_id"]] = _wl_entry(cand["paper_id"],
                                                  cand.get("evidence_sentence"))
            if len(out) >= k:
                break
    except Exception as exc:
        error = f"白名单构造失败：{type(exc).__name__}: {str(exc)[:120]}"
    for pid in (extra_ids or []):
        if pid in out or len(out) >= k + 4:
            continue
        ev = _best_evidence(pid, text)
        if ev:
            out[pid] = _wl_entry(pid, ev)
    return list(out.values()), error


def _best_evidence(paper_id: int, text: str) -> str | None:
    """对指定论文找一段能通过机械校验的证据句（向量句级优先，词面重叠兜底）。"""
    with db.conn() as c:
        row = c.execute("SELECT abstract FROM papers WHERE id=?", (paper_id,)).fetchone()
    if not row:
        return None
    abstract = row["abstract"] or ""
    try:
        import numpy as np
        tvec = np.asarray(embeddings.embed_texts([text[:2000]])[0], dtype=np.float32)
        ev, score = embeddings.best_sentence(tvec, paper_id)
        if ev and score >= 0.35 and cite._norm(ev) in cite._norm(abstract):
            return ev
    except Exception:
        pass
    ev, _ = cite._best_sentence_overlap(text, abstract)
    return ev if ev and cite._norm(ev) in cite._norm(abstract) else None


def _wl_entry(paper_id: int, evidence: str | None) -> dict:
    with db.conn() as c:
        r = c.execute("SELECT title, year, venue, authors, doi, arxiv_id "
                      "FROM papers WHERE id=?", (paper_id,)).fetchone()
    d = dict(r) if r else {"title": f"paper#{paper_id}"}
    d["paper_id"] = paper_id
    d["authors"] = json.loads(d.get("authors") or "[]")
    d["evidence_sentence"] = evidence or ""
    d["verified"] = True
    return d


# ── Agent 1：选题 ────────────────────────────────────────────────

def _topic_context(run: dict, limit: int = 20) -> tuple[list[dict], str]:
    """检索库内相关文献并组上下文。返回 (papers, context_text)。

    选题取 20 篇宽上下文，用迭代检索（派生查询只补位）：k≥10 实测有净增益、
    小 k 与单轮逐条相同、0 token——选题遗漏一篇关键文献比多看两篇噪声贵得多。
    """
    from . import deepsearch
    ids, _mode, _trace = deepsearch.deep_retrieve(run["topic"], limit)
    papers = []
    blocks = []
    with db.conn() as c:
        for pid in ids:
            r = c.execute("SELECT * FROM papers WHERE id=?", (pid,)).fetchone()
            if not r:
                continue
            try:
                card = json.loads(r["card_json"] or "{}")
            except json.JSONDecodeError:
                card = {}
            papers.append(dict(r) | {"card": card})
            lim = card.get("limitations") or ""
            parts = [f"[{r['id']}] {r['title']}（{r['venue'] or ''} {r['year'] or ''}）"]
            if card.get("tldr"):
                parts.append(f"    一句话：{card['tldr']}")
            if lim and lim != "（摘要未提及）":
                parts.append(f"    局限：{lim}")
            if r["abstract"]:
                parts.append(f"    摘要：{r['abstract'][:400]}")
            blocks.append("\n".join(parts))
    return papers, "\n\n".join(blocks)


def _as_list_of_dicts(value) -> list[dict]:
    """把模型返回的「一堆东西」规整成 list[dict]，规整不出来就给空列表。

    模型的形状是会漂的：同一个 prompt，这次给 `{"topics": [{...}]}`，
    下次可能给 `{"topics": 4}`（把数量当答案）、`{"topics": {...}}`（只给一个）、
    或者列表里混进字符串。而下游是 `for t in topics` + `len(t["evidence"])`——
    实测真模型上崩过一次：`TypeError: object of type 'int' has no len()`，
    烧掉 9657 token 之后整个选题阶段失败，而旁边就摆着写好的 `_offline_topics` 兜底。

    规整而不是抛异常，是因为调用方紧接着就有「空了就走离线」的分支——
    让形状问题落进那条既有的兜底路，比多一条异常路径简单。
    """
    if isinstance(value, dict):
        value = [value]
    if not isinstance(value, list):
        return []
    return [x for x in value if isinstance(x, dict)]


def stage_topic(run_id: str, progress=None) -> dict:
    """选题 Agent：库内缺口聚类 → 带证据的选题候选。产物 = topics_json，状态到检查点①。"""
    run = get_run(run_id)
    if not run:
        raise ValueError("写作任务不存在")
    if progress:
        progress(0.1, "topic", "库内检索相关文献")
    papers, context = _topic_context(run)
    if not papers:
        _update_run(run_id, status="failed", error="库内没有相关文献——先去「文献检索」采集")
        raise ValueError("库内没有相关文献——先去「文献检索」采集")

    topics = []
    degraded = None
    if llm.available():
        if progress:
            progress(0.4, "topic", f"基于 {len(papers)} 篇文献做缺口分析")
        user = (f"我的课题：{run['topic']}\n\n文献卡片（[id] 为文献编号，paper_id 用它）：\n{context}")
        raw = llm.chat(TOPIC_SYSTEM, user, purpose="write_topic", temperature=0.4)
        try:
            topics = _as_list_of_dicts(llm.extract_json(raw).get("topics"))
            if not topics:
                degraded = "选题输出的 topics 不是可用的对象列表，走离线缺口候选"
        except llm.LLMError as exc:
            degraded = f"选题输出未按 JSON 解析（{exc}），走离线缺口候选"
    if not topics:
        topics, degraded2 = _offline_topics(papers)
        degraded = degraded or degraded2

    # 机械核验：每条 evidence 的 quote 必须逐字回取到原文
    # 模型产物的后处理**整段兜住**。上面已经把 topics / evidence 规整成 list[dict]，
    # 但字段值的类型仍然不可信（title 是数字、quote 是数字、paper_id 是乱码…），
    # 而这一段每加一个字段就多一处可能崩的地方。与其逐个形状打补丁，
    # 不如让任何意外都落进下面既有的「空了就走离线」那条路——
    # 那条兜底就写在旁边，被一个 TypeError 炸掉等于白写。
    clean: list[dict] = []
    for t in topics:
        try:
            evs = _as_list_of_dicts(t.get("evidence"))
            for ev in evs:
                ev["verified"] = _verify_quote(ev.get("paper_id"), ev.get("quote"))
            t["evidence"] = evs
            t["title"] = str(t.get("title") or "").strip()
            t["evidence_ok"] = f"{sum(1 for e in evs if e['verified'])}/{len(evs)}"
            if t["title"]:
                clean.append(t)
        except Exception as exc:                        # noqa: BLE001
            degraded = degraded or f"有选题候选的字段形状异常已丢弃（{type(exc).__name__}）"
    if len(clean) < len(topics):
        degraded = degraded or f"{len(topics) - len(clean)} 个选题候选字段异常已丢弃"
    topics = clean
    if not topics:
        topics, degraded2 = _offline_topics(papers)
        degraded = degraded or degraded2
    if progress:
        progress(0.95, "topic", f"产出 {len(topics)} 个候选，等待人工定题")
    _update_run(run_id, topics=topics, status="pending_user_topic", stage="选题完成，等待定题",
                error=None)
    return {"topics": len(topics), "papers": len(papers), "degraded": degraded}


def _offline_topics(papers: list[dict]) -> tuple[list[dict], str]:
    """离线选题：从「局限」栏确定性抽取缺口候选（无 LLM 也能跑通流水线）。"""
    out = []
    for p in papers:
        lim = p["card"].get("limitations") or ""
        if not lim or lim == "（摘要未提及）" or len(lim) < 15:
            continue
        sent = next((s for s in embeddings.split_sentences(lim)), lim[:120])
        out.append({"title": f"从「{p['title'][:40]}」的局限出发：{run_topic_short(p)}",
                    "gap": sent, "angle": "（离线模式：直接以该文献自述局限为切入点）",
                    "risks": "（离线模式未做跨文献聚类，建议配置 LLM 后重跑选题）",
                    "evidence": [{"paper_id": p["id"], "quote": sent}],
                    "evidence_ok": "1/1"})
        if len(out) >= 3:
            break
    return out, "未配置 LLM：选题为库内局限栏的确定性抽取"


def run_topic_short(p: dict) -> str:
    t = p["card"].get("tldr") or p["title"]
    return t[:30]


# ── Agent 2：大纲 ────────────────────────────────────────────────

def select_topic(run_id: str, title: str) -> dict:
    """检查点①：人工定题（可改写候选题目）。只改状态，大纲在下一个任务里跑。"""
    run = get_run(run_id)
    if not run:
        raise ValueError("写作任务不存在")
    if run["status"] not in ("pending_user_topic", "outline_running", "failed"):
        raise ValueError(f"当前状态 {run['status']} 不允许定题")
    title = (title or "").strip()
    if len(title) < 4:
        raise ValueError("题目太短")
    _update_run(run_id, title=title, status="outline_running", stage="已定题，生成大纲中",
                error=None)
    return get_run(run_id)


def get_outline(run_id: str) -> dict | None:
    with db.conn() as c:
        r = c.execute("""SELECT * FROM outlines WHERE run_id=?
                         ORDER BY version DESC, id DESC LIMIT 1""", (run_id,)).fetchone()
    if not r:
        return None
    d = dict(r)
    d["sections"] = json.loads(d.pop("outline_json") or "[]")
    d["warnings"] = json.loads(d.pop("warnings_json") or "[]")
    return d


def validate_outline(sections: list[dict], valid_ids: set[int]) -> list[str]:
    """大纲机械校验：编号连续、字数预算为正、拟引文献必须在候选集内。"""
    issues = []
    for i, s in enumerate(sections, 1):
        no = int(s.get("no") or i)
        if no != i:
            issues.append(f"第 {i} 节编号不连续（no={no}）")
            s["no"] = i
        if not (s.get("title") or "").strip():
            issues.append(f"第 {i} 节缺标题")
        words = int(s.get("words") or 0)
        if not 100 <= words <= 3000:
            issues.append(f"第 {i} 节字数预算异常（{words}）")
            s["words"] = min(max(words, 200), 2000)
        bad = [int(x) for x in (s.get("paper_ids") or []) if int(x) not in valid_ids]
        if bad:
            issues.append(f"第 {i} 节拟引文献 {bad} 不在候选集内，已剔除")
            s["paper_ids"] = [int(x) for x in (s.get("paper_ids") or [])
                              if int(x) in valid_ids]
        if not s.get("paper_ids") and not s.get("no_support"):
            issues.append(f"第 {i} 节无拟引文献且未标 no_support，已补标（该节将如实「待补证据」）")
            s["no_support"] = True
    if not 2 <= len(sections) <= 8:
        issues.append(f"节数 {len(sections)} 超出 2-8 的范围")
    return issues


def stage_outline(run_id: str, progress=None) -> dict:
    """大纲 Agent：定题 + 文献卡片 → 大纲 JSON + 机械校验。产物 = outlines v1，状态到检查点②。"""
    run = get_run(run_id)
    if not run or not run["title"]:
        raise ValueError("还没定题")
    if progress:
        progress(0.1, "outline", "检索支撑文献")
    papers, _ = _topic_context(run, limit=16)
    valid_ids = {p["id"] for p in papers}
    card_ctx = "\n\n".join(
        f"[{p['id']}] {p['title']}（{p['year'] or ''}）\n"
        f"    一句话：{p['card'].get('tldr') or ''}\n"
        f"    方法：{(p['card'].get('method') or '')[:120]}\n"
        f"    结果：{(p['card'].get('results') or '')[:120]}"
        for p in papers)

    outline, warnings, degraded = None, [], None
    if llm.available():
        if progress:
            progress(0.35, "outline", "生成结构化大纲")
        user = (f"题目：{run['title']}\n课题：{run['topic']}\n目标总字数："
                f"{run['config'].get('target_words', 3000)}\n\n文献卡片：\n{card_ctx}")
        raw = llm.chat(OUTLINE_SYSTEM, user, purpose="write_outline", temperature=0.35)
        try:
            data = llm.extract_json(raw)
            # sections 也会漂形状；不规整的话 `validate_outline` 与
            # `len(outline["sections"])` 会拿到 int 直接崩。
            outline = {"title": data.get("title") or run["title"],
                       "sections": _as_list_of_dicts(data.get("sections"))}
        except llm.LLMError as exc:
            degraded = f"大纲输出未按 JSON 解析（{exc}），走离线模板"
    if not outline or not outline.get("sections"):
        outline = {"title": run["title"], "sections": _offline_outline(papers)}
        degraded = degraded or "未配置 LLM：大纲为离线模板"

    warnings = validate_outline(outline["sections"], valid_ids)
    if progress:
        progress(0.9, "outline", f"大纲 {len(outline['sections'])} 节，"
                                 f"校验{'通过' if not warnings else f' {len(warnings)} 条备注'}")
    # 失败重跑 / 重新生成时版本递增，不覆盖旧版（可回滚对比）
    with db.conn() as c:
        prev = c.execute("SELECT COALESCE(MAX(version),0) v FROM outlines WHERE run_id=?",
                         (run_id,)).fetchone()["v"]
        c.execute("""INSERT INTO outlines(run_id,version,title,outline_json,warnings_json)
                     VALUES(?,?,?,?,?)""",
                  (run_id, prev + 1, outline["title"],
                   json.dumps(outline["sections"], ensure_ascii=False),
                   json.dumps(warnings, ensure_ascii=False)))
        c.commit()
    glossary = _build_glossary(papers)
    _update_run(run_id, status="pending_user_outline", stage="大纲完成，等待确认或修改",
                glossary=glossary, error=None)
    return {"sections": len(outline["sections"]), "warnings": warnings, "degraded": degraded}


def _offline_outline(papers: list[dict]) -> list[dict]:
    ids = [p["id"] for p in papers]
    return [
        {"no": 1, "title": "研究背景与问题", "points": ["说明研究场景与待解决问题"],
         "paper_ids": ids[:2], "words": 400},
        {"no": 2, "title": "相关工作", "points": ["按方法路线比较代表工作，给出引用"],
         "paper_ids": ids[2:8], "words": 900},
        {"no": 3, "title": "方法与分析框架", "points": ["归纳主流方法与评价指标"],
         "paper_ids": ids[8:12], "words": 800},
        {"no": 4, "title": "开放问题与展望", "points": ["指出库内文献未覆盖的方向"],
         "paper_ids": ids[12:14], "words": 500},
    ]


def _build_glossary(papers: list[dict]) -> dict[str, str]:
    """从卡片 keywords 聚合高频术语（长文一致性用，值 = 出处论文数）。"""
    freq: dict[str, int] = {}
    for p in papers:
        for kw in p["card"].get("keywords") or []:
            kw = str(kw).strip()
            if 2 <= len(kw) <= 30:
                freq[kw] = freq.get(kw, 0) + 1
    top = sorted(freq, key=lambda k: -freq[k])[:10]
    return {k: f"{freq[k]} 篇文献的关键词" for k in top}


def save_outline(run_id: str, sections: list[dict] | None) -> dict:
    """检查点②：人工改纲（sections=None 表示不改直接开写）。返回更新后的 run。"""
    run = get_run(run_id)
    if not run:
        raise ValueError("写作任务不存在")
    cur = get_outline(run_id)
    if not cur:
        raise ValueError("大纲还没生成")
    if run["status"] not in ("pending_user_outline", "drafting", "drafted", "failed"):
        raise ValueError(f"当前状态 {run['status']} 不允许确认大纲")
    if sections is None:
        version, edited, secs = cur["version"], 0, cur["sections"]
    else:
        sections = sorted(sections, key=lambda s: int(s.get("no") or 0))
        version = cur["version"] + 1
        edited = 1
        secs = sections
        with db.conn() as c:
            c.execute("""INSERT INTO outlines(run_id,version,title,outline_json,
                         warnings_json,edited) VALUES(?,?,?,?,?,1)""",
                      (run_id, version, cur["title"],
                       json.dumps(secs, ensure_ascii=False), "[]"))
            c.commit()
    _update_run(run_id, status="drafting", stage=f"按大纲 v{version} 逐节撰写中", error=None)
    return get_run(run_id)


# ── Agent 3：撰写（含节评审闭环）──────────────────────────────────

def _cache_key(title: str, points: list, words: int, glossary: dict) -> str:
    material = json.dumps([PROMPT_VER, title, points, words,
                           sorted(glossary.items())], ensure_ascii=False, sort_keys=True)
    return hashlib.sha1(material.encode("utf-8")).hexdigest()[:16]


def _find_cached(run_id: str, cache_key: str) -> dict | None:
    with db.conn() as c:
        r = c.execute("""SELECT * FROM sections WHERE run_id=? AND cache_key=?
                         AND status IN ('done','reused') AND content != ''
                         ORDER BY id DESC LIMIT 1""", (run_id, cache_key)).fetchone()
    return dict(r) if r else None


def _upsert_section(run_id: str, outline_version: int, sec: dict, **fields):
    base = {"run_id": run_id, "outline_version": outline_version,
            "sec_no": int(sec["no"]), "title": sec["title"],
            "points_json": json.dumps(sec.get("points") or [], ensure_ascii=False)}
    sets = list(base.items()) + [(k, json.dumps(v, ensure_ascii=False)
                                  if isinstance(v, (dict, list)) else v)
                                 for k, v in fields.items()]
    cols = ", ".join(k for k, _ in sets)
    ph = ", ".join("?" for _ in sets)
    upd = ", ".join(f"{k}=excluded.{k}" for k, _ in sets
                    if k not in ("run_id", "outline_version", "sec_no"))
    with db.conn() as c:
        c.execute(f"""INSERT INTO sections({cols}) VALUES({ph})
                      ON CONFLICT(run_id,outline_version,sec_no) DO UPDATE SET {upd},
                      updated_at=datetime('now','localtime')""", tuple(v for _, v in sets))
        c.commit()


def _latest_sections(run_id: str) -> list[dict]:
    out = get_outline(run_id)
    if not out:
        return []
    with db.conn() as c:
        rows = c.execute("""SELECT * FROM sections WHERE run_id=? AND outline_version=?
                            ORDER BY sec_no""", (run_id, out["version"])).fetchall()
    return [dict(r) for r in rows]


def _draft_section(run: dict, sec: dict, whitelist: list[dict], attempt: int,
                   feedback: list[dict] | None) -> str:
    """撰写 Agent：一节草稿。attempt=1 初稿；>1 带评审意见重写。"""
    context, _sources = writing._paper_context([w["paper_id"] for w in whitelist],
                                               sec["title"] + " " + " ".join(sec.get("points") or []))
    other_titles = "；".join(s["title"] for s in
                             (get_outline(run["id"]) or {}).get("sections", [])
                             if s["no"] != sec["no"])
    user = (f"章节任务：第 {sec['no']} 节「{sec['title']}」\n"
            f"论证要点（必须覆盖）：\n- " + "\n- ".join(sec.get("points") or ["围绕节标题展开"]) + "\n"
            f"引用白名单（[n] 只能指这些文献，n 为下列序号）：\n{context or '（白名单为空：本节不要出现任何引用编号）'}\n"
            f"术语表（用词须一致）：{'、'.join(run['glossary']) or '（无）'}\n"
            f"其它章节标题（衔接参考）：{other_titles}\n"
            f"本节目标字数：约 {sec.get('words', 600)} 字")
    if attempt > 1 and feedback:
        lines = "\n".join(f"- [{it.get('severity')}] {it.get('location')}：{it.get('problem')}→{it.get('fix')}"
                          for it in feedback[:6])
        user += f"\n\n上一稿评审意见（逐条解决，保持事实与引用编号不变）：\n{lines}"
    purpose = "write_draft" if attempt == 1 else f"write_rewrite{attempt}"
    return llm.chat(WRITER_SYSTEM, user, purpose=purpose, temperature=0.3).strip()


def _sanitize(text: str, whitelist: list[dict]) -> tuple[str, list[int]]:
    """白名单外编号剥离（如实记录，不静默）：返回 (清洗后文本, 被剥离的编号)。"""
    n = len(whitelist)
    removed = []
    def repl(m):
        i = int(m.group(1))
        if 1 <= i <= n:
            return m.group(0)
        removed.append(i)
        return ""
    return re.sub(r"\[(\d{1,2})\]", repl, text), sorted(set(removed))


def _offline_draft(sec: dict, whitelist: list[dict]) -> str:
    """离线模式：确定性结构占位稿（不假装是 LLM 产物），整条流水线无 key 可跑通。"""
    pts = sec.get("points") or ["围绕节标题展开论述"]
    lines = [f"- {p}。（待补证据：离线模式未接 LLM，配置 .env 后重新生成本节。）" for p in pts]
    if whitelist:
        lines.append(f"库内已有 {len(whitelist)} 篇通过证据核验的文献可支撑本节"
                     f"（如 [{1}]），具体结论需回原文核对后展开。")
    else:
        lines.append("库内暂无通过证据句核验的文献支撑本节——这是如实的「无支撑段」，不编造引用。")
    return "\n\n".join(lines)


def _write_one(run: dict, sec: dict, outline_version: int, force: bool,
               progress=None, frac0: float = 0.0, frac1: float = 1.0) -> dict:
    """单节全流程：缓存查 → 白名单 → 撰写⇄评审重写闭环 → 落库。"""
    run_id = run["id"]
    key = _cache_key(sec["title"], sec.get("points") or [], sec.get("words", 600),
                     run["glossary"])
    if not force:
        hit = _find_cached(run_id, key)
        if hit:
            _upsert_section(run_id, outline_version, sec,
                            whitelist_json=json.loads(hit["whitelist_json"] or "[]"),
                            cache_key=key, content=hit["content"],
                            citations_json=hit["citations_json"],
                            removed_json=hit["removed_json"], status="reused",
                            attempt=hit["attempt"], score=hit["score"],
                            review_json=hit["review_json"])
            if progress:
                progress(frac1, f"节{sec['no']}", f"「{sec['title'][:20]}」缓存命中，0 token 复用")
            return {"sec_no": sec["no"], "status": "reused", "cache_hit": True}

    if progress:
        progress(frac0 + 0.02, f"节{sec['no']}", f"「{sec['title'][:20]}」文献 Agent 核验白名单")
    wl, wl_error = _whitelist(sec["title"] + " " + " ".join(sec.get("points") or []),
                              k=6, extra_ids=sec.get("paper_ids") or [])

    if not llm.available():
        # 离线：结构占位稿 + 真实白名单（有则挂），不调评审、不盲目重写
        content, removed = _offline_draft(sec, wl), []
        citations = {str(i + 1): wl[i] for i in range(len(wl))
                     if f"[{i + 1}]" in content}
        _upsert_section(run_id, outline_version, sec, whitelist_json=wl, cache_key=key,
                        content=content, citations_json=citations, removed_json=removed,
                        status="done", attempt=1, score=None, review_json=None)
        if progress:
            progress(frac1, f"节{sec['no']}", f"「{sec['title'][:20]}」离线模板稿（未接 LLM）")
        return {"sec_no": sec["no"], "status": "done", "attempt": 1, "score": None,
                "citations": len(citations), "removed": 0, "cache_hit": False,
                "degraded": f"offline-template；{wl_error}" if wl_error
                            else "offline-template"}

    max_rewrites = int(run["config"].get("max_rewrites", 2))
    attempt, feedback = 0, None
    content, citations, removed, review = "", {}, [], None
    met_bar = False        # 是否达到评审门槛（用尽重写次数仍不达标时如实交付并标注）
    while attempt <= max_rewrites:
        attempt += 1
        if progress:
            progress(frac0 + 0.1 + (0.5 / (max_rewrites + 1)) * (attempt - 1),
                     f"节{sec['no']}", f"「{sec['title'][:20]}」撰写（第 {attempt} 稿）")
        content = _draft_section(run, sec, wl, attempt, feedback)
        content, removed = _sanitize(content, wl)
        citations = {str(i + 1): wl[i] for i in range(len(wl))
                     if f"[{i + 1}]" in content}
        if progress:
            progress(frac0 + 0.1 + (0.5 / (max_rewrites + 1)) * (attempt - 1) + 0.2,
                     f"节{sec['no']}", f"「{sec['title'][:20]}」评审（重档模型）")
        review = writing.review_text(content, paper_ids=[w["paper_id"] for w in wl])
        score = review.get("score")
        majors = [i for i in (review.get("issues") or []) if i.get("severity") == "major"]
        if score is not None and score >= CHECKPOINT_THRESHOLD and not majors:
            met_bar = True
            break  # 达标
        if score is None and not majors:
            break  # 离线/解析失败：无从判定，不盲目重写
        feedback = (review.get("issues") or [])[:6]
        # 不达标 → 带意见重写；用尽重写次数则如实交付（score/review 落库，UI 明示）
    _upsert_section(run_id, outline_version, sec,
                    whitelist_json=wl, cache_key=key, content=content,
                    citations_json=citations, removed_json=removed,
                    status="done", attempt=attempt,
                    score=review.get("score") if review else None,
                    review_json=review)
    return {"sec_no": sec["no"], "status": "done", "attempt": attempt,
            "met_bar": met_bar,
            "score": review.get("score") if review else None,
            "citations": len(citations), "removed": len(removed),
            "cache_hit": False, "degraded": wl_error}


def stage_sections(run_id: str, progress=None, force_sec: int | None = None) -> dict:
    """撰写阶段：逐节「撰写⇄文献白名单⇄评审重写」。force_sec 指定则只重跑该节（绕缓存）。"""
    run = get_run(run_id)
    outline = get_outline(run_id)
    if not run or not outline:
        raise ValueError("大纲不存在")
    secs = outline["sections"]
    _update_run(run_id, status="drafting", error=None)
    results = []
    n = len(secs)
    for i, sec in enumerate(secs):
        f0, f1 = 0.05 + i / n * 0.9, 0.05 + (i + 1) / n * 0.9
        if force_sec is not None and int(sec["no"]) != int(force_sec):
            continue
        results.append(_write_one(run, sec, outline["version"],
                                  force=(force_sec is not None),
                                  progress=progress, frac0=f0, frac1=f1))
    rows = _latest_sections(run_id)
    failed = [r["sec_no"] for r in rows if r["status"] == "failed"]
    reused = sum(1 for r in results if r.get("cache_hit"))
    # 用尽重写次数仍未达评审门槛的节：如实交付但单列出来，不假装全绿
    below_bar = [r["sec_no"] for r in results
                 if r.get("score") is not None and not r.get("met_bar")]
    _update_run(run_id, status="drafted" if not failed else "failed",
                stage="全节完成" if not failed else f"第 {failed} 节失败，可单节重跑",
                error=None)
    return {"sections": len(results), "cache_reused": reused,
            "failed": failed, "below_bar": below_bar,
            "scores": {r["sec_no"]: r.get("score") for r in results}}


# ── 装配与机械终检 ───────────────────────────────────────────────

def assemble(run_id: str) -> tuple[str, dict[int, dict], list[dict]]:
    """把各节装配成全文：局部 [n] → 全局编号重排，返回 (正文, 全局refs, 各节信息)。"""
    outline = get_outline(run_id)
    refs: dict[int, dict] = {}
    gmap: dict[int, int] = {}  # paper_id → 全局编号
    infos = []
    parts = []
    for row in _latest_sections(run_id):
        cit = json.loads(row["citations_json"] or "{}")
        text = row["content"] or ""

        def repl(m):
            entry = cit.get(m.group(1))
            if not entry:
                return m.group(0)  # 悬空编号保留，终检如实报告
            pid = entry["paper_id"]
            if pid not in gmap:
                gmap[pid] = len(gmap) + 1
                refs[gmap[pid]] = entry
            return f"[{gmap[pid]}]"

        text2 = re.sub(r"\[(\d{1,2})\]", repl, text)
        infos.append({"sec_no": row["sec_no"], "title": row["title"],
                      "status": row["status"], "score": row["score"],
                      "attempt": row["attempt"],
                      "removed": json.loads(row["removed_json"] or "[]")})
        parts.append(f"## {row['sec_no']}. {row['title']}\n\n{text2}")
    title = (outline or {}).get("title") or ""
    body = (f"# {title}\n\n" + "\n\n".join(parts)) if title else "\n\n".join(parts)
    return body, refs, infos


def _reference_lines(refs: dict[int, dict]) -> list[str]:
    out = []
    for no in sorted(refs):
        e = refs[no]
        with db.conn() as c:
            r = c.execute("SELECT title, authors, year, venue, doi, arxiv_id "
                          "FROM papers WHERE id=?", (e["paper_id"],)).fetchone()
        if r:
            out.append(f"[{no}] {cite.to_gbt(dict(r))}")
        else:
            out.append(f"[{no}] {e.get('title', '')}")
    return out


def final_check(text: str, refs: dict[int, dict]) -> dict:
    """机械终检（不加 LLM）：悬空/未核验编号、待补证据数、字数、引用核验率。"""
    used = [int(x) for x in re.findall(r"\[(\d{1,2})\]", text)]
    dangling = sorted({n for n in used if n not in refs})
    unused = sorted(n for n in refs if n not in used)
    unverified = sorted(n for n, r in refs.items() if not r.get("verified"))
    verified_used = [n for n in used if n in refs and refs[n].get("verified")]
    rate = round(len(verified_used) / len(used), 3) if used else None
    clean = re.sub(r"\[\d{1,2}\]", "", text)
    return {"ok": not dangling and not unverified,
            "citation_count": len(used),
            "verified_rate": rate,               # 草稿引用核验率
            "dangling": dangling, "unused": unused, "unverified": unverified,
            "pending_evidence": clean.count("待补证据"),
            "words": len(re.sub(r"\s+", "", clean))}


def report(run_id: str) -> dict:
    """任意时刻可算的引用一致性报告（纯机械，不加 LLM）：装配 → 终检。"""
    body, refs, infos = assemble(run_id)
    rep = final_check(body, refs)
    rep["sections"] = infos
    rep["references"] = [r.split("] ", 1)[-1] for r in _reference_lines(refs)]
    return {"body": body, "report": rep}


def _export_dir():
    d = config.DATA_DIR / "exports" / "write"
    d.mkdir(parents=True, exist_ok=True)
    return d


def _md_document(run: dict, body: str, refs: dict[int, dict], report: dict) -> str:
    lines = [body, ""]
    if refs:
        lines += ["## 参考文献", ""] + _reference_lines(refs) + [""]
    else:
        lines += ["## 参考文献", "",
                  "（本稿暂无引用——库内没有能通过证据句核验的支撑文献，按铁律不硬凑引用。）", ""]
    rate = report.get("verified_rate")
    rate_s = f"{rate * 100:.1f}%" if isinstance(rate, float) else "—"
    lines.append(f"> **AI 参与说明**：本稿草稿由 PaperNest 多 Agent 写作流水线生成，"
                 f"定题与改纲两个人工检查点由作者确认。引用 {report.get('citation_count', 0)} 条"
                 f"均来自库内真实文献并经证据句机械回取校验（核验率 {rate_s}）；"
                 f"本稿不生成实验数据与结果，「待补证据」标记处需作者人工补充。")
    return "\n".join(lines)


def _docx_document(path, run: dict, body: str, refs: dict[int, dict], report: dict):
    from docx import Document
    doc = Document()
    md = _md_document(run, body, refs, report)
    for raw in md.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        clean = re.sub(r"\*\*|``?|__", "", line)
        if line.startswith("### "):
            doc.add_heading(clean[4:], level=2)
        elif line.startswith("## "):
            doc.add_heading(clean[3:], level=1)
        elif line.startswith("# "):
            doc.add_heading(clean[2:], level=0)
        elif line.startswith("> "):
            doc.add_paragraph(clean[2:], style="Intense Quote")
        elif line.startswith("- "):
            doc.add_paragraph(clean[2:], style="List Bullet")
        else:
            doc.add_paragraph(clean)
    doc.save(str(path))


# ── Agent 4：润色 + 终检 + 导出 ──────────────────────────────────

def stage_polish(run_id: str, progress=None) -> dict:
    run = get_run(run_id)
    rows = _latest_sections(run_id)
    if not run or not rows:
        raise ValueError("还没有可装配的章节")
    if progress:
        progress(0.1, "assemble", "装配全文（局部 [n] → 全局编号）")
    body, refs, infos = assemble(run_id)
    removed_total = sum(len(i["removed"]) for i in infos)
    if progress:
        progress(0.35, "polish", "润色 Agent（重档模型，保留事实与引用编号）")
    polished = writing.polish_text(
        body, instruction="统一术语与行文风格，保持所有事实、数字与 [n] 引用编号完全不变",
        paper_ids=[e["paper_id"] for e in refs.values()])
    text = polished.get("revised") or body
    if progress:
        progress(0.8, "final-check", "机械终检（悬空/未核验编号、待补证据、字数）")
    report = final_check(text, refs)
    report["polish_degraded"] = polished.get("degraded")
    report["removed_local_citations"] = removed_total
    report["sections"] = infos
    report["references"] = [r.split("] ", 1)[-1] for r in _reference_lines(refs)]

    md_path = _export_dir() / f"{run_id}.md"
    md_path.write_text(_md_document(run, text, refs, report), encoding="utf-8")
    files = {"md": str(md_path)}
    try:
        docx_path = _export_dir() / f"{run_id}.docx"
        _docx_document(docx_path, run, text, refs, report)
        files["docx"] = str(docx_path)
    except ImportError:
        report["docx_note"] = "未安装 python-docx，仅导出 Markdown"
    report["files"] = files
    if progress:
        progress(0.98, "done", f"完成：{report['words']} 字，引用核验率 "
                               f"{report['verified_rate'] if report['verified_rate'] is not None else '—'}")
    _update_run(run_id, status="done", stage="终稿已导出", result=report, error=None)
    return {"words": report["words"], "verified_rate": report["verified_rate"],
            "ok": report["ok"], "files": files}


# ── 任务分发（jobs.kind='write'）─────────────────────────────────

def attach_job(run_id: str, job_id: str):
    """把流水线 run 与正在跑的异步任务关联（UI 从 run 一眼看到任务 id）。"""
    _update_run(run_id, job_id=job_id)


def run_stage(run_id: str, params: dict, progress=None) -> dict:
    stage = params.get("stage", "")
    run = get_run(run_id)
    if not run:
        raise ValueError("写作任务不存在")
    try:
        if stage == "topic":
            return stage_topic(run_id, progress)
        if stage == "outline":
            return stage_outline(run_id, progress)
        if stage == "sections":
            return stage_sections(run_id, progress)
        if stage == "section":
            return stage_sections(run_id, progress, force_sec=int(params["sec_no"]))
        if stage == "polish":
            return stage_polish(run_id, progress)
        raise ValueError(f"未知写作阶段：{stage}")
    except Exception as exc:
        _update_run(run_id, status="failed", error=f"{type(exc).__name__}: {str(exc)[:300]}")
        raise


def get_run_state(run_id: str) -> dict | None:
    """前端一次拉全：run + 大纲（最新版）+ 各节状态/引用 + 终检报告。"""
    run = get_run(run_id)
    if not run:
        return None
    outline = get_outline(run_id)
    sections = []
    if outline:
        for r in _latest_sections(run_id):
            d = dict(r)
            d["points"] = json.loads(d.pop("points_json") or "[]")
            d["whitelist"] = json.loads(d.pop("whitelist_json") or "[]")
            d["citations"] = json.loads(d.pop("citations_json") or "{}")
            d["removed"] = json.loads(d.pop("removed_json") or "[]")
            d["review"] = json.loads(d.pop("review_json")) if d.pop("review_json") else None
            sections.append(d)
    return {"run": run, "outline": outline, "sections": sections}
