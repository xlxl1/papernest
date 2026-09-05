"""多格式文档解析的离线测试。

所有样本文件都在测试里**程序化生成**（python-docx 造 docx、python-pptx 造 pptx、
手写 html/md/txt 字符串落盘），不依赖任何外部样本，因此跨机器可复现。
本模块不联网、不调 LLM、不写库，但仍按项目约定把 DB_PATH/DATA_DIR 指向临时目录、
把 llm/embeddings 摁成不可用——将来有人往 docparse 里加了这类调用，测试必须立刻报警。
"""
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from pptx import Presentation
from pptx.util import Inches

from papernest import config, docparse

# ── 样本内容（写死，断言可以手算） ──

MD_TEXT = """---
title: 近场信道估计笔记
tags: 组会
---

# 近场信道估计

正文第一段，讨论球面波模型。
同一段的第二行。

## 实验设置

```python
# 这一行是代码注释，绝不能被当成一级标题
def loss(x):
    return x ** 2

# 空行也不能把代码块切碎
```

| 模型 | 准确率 | F1 |
| --- | --- | --- |
| BERT | 0.83 | 0.81 |
| Ours | 0.91 | 0.90 |

结尾段落。
"""

HTML_TEXT = """<!DOCTYPE html>
<html><head>
<title>Near-field &amp; XL-MIMO 综述</title>
<meta name="description" content="一篇网页存档">
<style>body { color: #333; }</style>
<script>var tracker = 1; if (a < b) { alert("x"); }</script>
</head>
<body>
<nav>首页 登录 关于我们</nav>
<h1>近场信道估计</h1>
<p>实体测试：AT&amp;T 与 &lt;tag&gt; 以及 &#39;引号&#39;。</p>
<h2>方法</h2>
<p>这里有个没闭合的段落
<div>紧跟着一个 div
<table>
<tr><th>模型</th><th>准确率</th></tr>
<tr><td>BERT</td><td>0.83</td></tr>
<tr><td>Ours</td><td>0.91</td>
</table>
<footer>版权所有 2026</footer>
"""

TXT_TEXT = """第一段中文，用于验证 GBK 解码。这里放一些常见汉字：矩阵、信道、稀疏、重构。

第二段：空行分段。

第三段结束。
"""


def make_docx(path: Path) -> str:
    """一份「像投稿稿」的 docx：标题层级 + 中文正文 + 夹在段落之间的三行表格。"""
    doc = Document()
    doc.core_properties.title = "多格式解析测试文档"
    doc.core_properties.author = "张三"
    doc.add_heading("绪论", level=1)
    doc.add_paragraph("这是中文正文段落，用于验证编码与顺序。")
    doc.add_heading("方法", level=2)
    doc.add_paragraph("方法段落在表格之前。")
    t = doc.add_table(rows=3, cols=3)
    for i, row in enumerate((("模型", "准确率", "F1"),
                            ("BERT", "0.83", "0.81"),
                            ("Ours", "0.91", "0.90"))):
        for j, v in enumerate(row):
            t.cell(i, j).text = v
    doc.add_paragraph("表格之后的段落。")
    doc.add_heading("小结", level=3)
    doc.add_paragraph("")          # 空段落不该产生空块
    doc.save(str(path))
    return str(path)


def make_pptx(path: Path) -> str:
    """4 页：正常页 + 带表格页 + 空白页 + 只有标题页；第 1 页带演讲者备注。"""
    prs = Presentation()
    prs.core_properties.title = "组会汇报"

    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "研究背景"
    s1.placeholders[1].text_frame.text = "近场区域的球面波模型\n远场近似在大孔径下失效"
    s1.notes_slide.notes_text_frame.text = "这里解释为什么远场近似会失效，属于讲稿内容。"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])   # Title Only
    s2.shapes.title.text = "实验结果"
    tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2),
                              Inches(4), Inches(1)).table
    tbl.cell(0, 0).text = "方法"
    tbl.cell(0, 1).text = "NMSE"
    tbl.cell(1, 0).text = "Ours"
    tbl.cell(1, 1).text = "-21.3"

    prs.slides.add_slide(prs.slide_layouts[6])        # 完全空白页

    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "结论"

    prs.save(str(path))
    return str(path)


def make_broken_docx(path: Path) -> str:
    """魔数与部件都对得上（会被判成 word），但 XML 是垃圾——用来制造 failed。"""
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("word/document.xml", "这根本不是 XML")
    return str(path)


class _Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="pn_docparse_"))
        self.addCleanup(shutil.rmtree, self.tmp, True)
        self._old_db, self._old_data = config.DB_PATH, config.DATA_DIR
        config.DATA_DIR = self.tmp / "data"
        config.DB_PATH = config.DATA_DIR / "t.db"
        self._patches = [
            mock.patch("papernest.embeddings.available", return_value=False),
            mock.patch("papernest.llm.available", return_value=False),
        ]
        for p in self._patches:
            p.start()

    def tearDown(self):
        for p in self._patches:
            p.stop()
        config.DB_PATH, config.DATA_DIR = self._old_db, self._old_data
        shutil.rmtree(self.tmp, ignore_errors=True)

    def write(self, name: str, text: str, encoding: str = "utf-8") -> Path:
        p = self.tmp / name
        p.write_bytes(text.encode(encoding))
        return p

    @staticmethod
    def kinds(res: dict) -> list[str]:
        return [b["kind"] for b in res["blocks"]]

    @staticmethod
    def of_kind(res: dict, kind: str) -> list[dict]:
        return [b for b in res["blocks"] if b["kind"] == kind]


# ── Word ──

class TestDocx(_Base):
    def test_heading_levels_and_order(self):
        res = docparse.extract_docx(make_docx(self.tmp / "a.docx"))
        heads = [(b["text"], b["level"]) for b in self.of_kind(res, "heading")]
        self.assertEqual(heads, [("绪论", 1), ("方法", 2), ("小结", 3)])
        # 表格必须夹在「方法段落」和「表格之后的段落」之间，顺序不能被打乱
        self.assertEqual(self.kinds(res),
                         ["heading", "paragraph", "heading", "paragraph",
                          "table", "paragraph", "heading"])

    def test_table_kept_whole_with_header(self):
        res = docparse.extract_docx(make_docx(self.tmp / "a.docx"))
        tables = self.of_kind(res, "table")
        self.assertEqual(len(tables), 1, "一张表必须恒定是一个块，不能被打散")
        t = tables[0]
        self.assertEqual(t["rows"][0], ["模型", "准确率", "F1"], "表头行必须保留")
        self.assertEqual(t["rows"][2], ["Ours", "0.91", "0.90"])
        self.assertEqual(t["text"].split("\n")[0], "模型 | 准确率 | F1")
        self.assertEqual((t["n_rows"], t["n_cols"]), (3, 3))

    def test_chinese_and_meta(self):
        res = docparse.extract_docx(make_docx(self.tmp / "a.docx"))
        self.assertIn("这是中文正文段落，用于验证编码与顺序。", res["text"])
        self.assertEqual(res["title"], "多格式解析测试文档")
        self.assertEqual(res["meta"]["author"], "张三")
        self.assertIn("created", res["meta"])
        self.assertEqual(res["warnings"], [])

    def test_no_empty_blocks(self):
        res = docparse.extract_docx(make_docx(self.tmp / "a.docx"))
        self.assertTrue(all(b["text"].strip() for b in res["blocks"]))

    def test_title_falls_back_to_first_heading_not_body(self):
        doc = Document()
        doc.add_paragraph("正文首行绝不能被当成标题")
        doc.add_heading("真正的标题", level=1)
        doc.save(str(self.tmp / "b.docx"))
        res = docparse.extract_docx(self.tmp / "b.docx")
        self.assertEqual(res["title"], "真正的标题")

    def test_title_none_when_nothing_to_take(self):
        doc = Document()
        doc.add_paragraph("只有正文，没有标题，也没有文档属性")
        doc.save(str(self.tmp / "c.docx"))
        self.assertIsNone(docparse.extract_docx(self.tmp / "c.docx")["title"])


# ── PowerPoint ──

class TestPptx(_Base):
    def setUp(self):
        super().setUp()
        self.res = docparse.extract_pptx(make_pptx(self.tmp / "a.pptx"))

    def test_one_block_per_slide_indexed_by_page(self):
        slides = self.of_kind(self.res, "slide")
        # 4 页里第 3 页是空白页，不产生任何块
        self.assertEqual([b["index"] for b in slides], [1, 2, 4])
        self.assertEqual(self.res["meta"]["n_slides"], 4)

    def test_blank_slide_produces_nothing(self):
        self.assertEqual([b for b in self.res["blocks"] if b["index"] == 3], [])

    def test_title_placeholder_marked(self):
        s1 = self.of_kind(self.res, "slide")[0]
        self.assertEqual(s1["title"], "研究背景")
        self.assertTrue(s1["has_title"])
        # slide 块是自包含的一页：标题行 + 正文都在里面
        self.assertTrue(s1["text"].startswith("研究背景"))
        self.assertIn("远场近似在大孔径下失效", s1["text"])

    def test_notes_collected(self):
        notes = self.of_kind(self.res, "notes")
        self.assertEqual(len(notes), 1)
        self.assertEqual(notes[0]["index"], 1)
        self.assertIn("为什么远场近似会失效", notes[0]["text"])

    def test_table_on_slide_kept_whole(self):
        tables = self.of_kind(self.res, "table")
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0]["index"], 2, "表格块也要能定位到页码")
        self.assertEqual(tables[0]["rows"], [["方法", "NMSE"], ["Ours", "-21.3"]])

    def test_title_from_core_properties(self):
        self.assertEqual(self.res["title"], "组会汇报")


# ── HTML ──

class TestHtml(_Base):
    def setUp(self):
        super().setUp()
        self.res = docparse.extract_html(HTML_TEXT)

    def test_script_style_nav_footer_dropped(self):
        text = self.res["text"]
        for junk in ("tracker", "color: #333", "首页 登录 关于我们", "版权所有"):
            self.assertNotIn(junk, text, f"{junk} 应该被丢掉")

    def test_entities_unescaped_once(self):
        para = self.of_kind(self.res, "paragraph")[0]["text"]
        self.assertEqual(para, "实体测试：AT&T 与 <tag> 以及 '引号'。")
        self.assertEqual(self.res["title"], "Near-field & XL-MIMO 综述")

    def test_headings_with_level(self):
        self.assertEqual([(b["text"], b["level"]) for b in self.of_kind(self.res, "heading")],
                         [("近场信道估计", 1), ("方法", 2)])

    def test_table_rows_preserved_even_unclosed(self):
        tables = self.of_kind(self.res, "table")
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0]["rows"],
                         [["模型", "准确率"], ["BERT", "0.83"], ["Ours", "0.91"]])

    def test_unclosed_tags_do_not_crash(self):
        """</h1> </p> </td> 全没有——手写/裁剪过的页面里这是常态。"""
        broken = "<html><body><h1>标题<p>段落<table><tr><td>a<td>b</table><div>尾巴"
        res = docparse.extract_html(broken)
        self.assertEqual([b["kind"] for b in res["blocks"]],
                         ["heading", "paragraph", "table", "paragraph"])
        self.assertEqual(res["blocks"][0]["text"], "标题",
                         "漏了 </h1> 不能把后面整篇正文并进标题块")
        self.assertEqual(self.of_kind(res, "table")[0]["rows"], [["a", "b"]])

    def test_unclosed_table_still_emitted_with_warning(self):
        res = docparse.extract_html("<body><table><tr><td>x</td><td>y</td></tr>")
        self.assertEqual(self.of_kind(res, "table")[0]["rows"], [["x", "y"]])
        self.assertTrue(any("未闭合" in w for w in res["warnings"]))

    def test_unclosed_skip_tag_does_not_eat_the_page(self):
        """<nav> 忘了闭合是常见的手写页面缺陷，不能把整篇正文吞掉。"""
        res = docparse.extract_html(
            "<html><body><nav>菜单 登录<h1>正文标题</h1><p>正文内容</p></body></html>")
        self.assertIn("正文内容", res["text"])
        self.assertIn("正文标题", res["text"])
        self.assertNotIn("菜单", res["text"], "nav 里的样板文字仍然要丢掉")
        self.assertTrue(any("强制结束跳过区" in w for w in res["warnings"]))

    def test_meta_description(self):
        self.assertEqual(self.res["meta"]["description"], "一篇网页存档")

    def test_accepts_file_path(self):
        p = self.write("page.html", HTML_TEXT)
        self.assertEqual(docparse.extract_html(str(p))["title"],
                         "Near-field & XL-MIMO 综述")

    def test_pre_kept_as_code(self):
        res = docparse.extract_html("<body><pre>def f():\n    return 1</pre></body>")
        code = self.of_kind(res, "code")
        self.assertEqual(len(code), 1)
        self.assertIn("    return 1", code[0]["text"], "pre 里的缩进不能被压掉")


# ── Markdown ──

class TestMarkdown(_Base):
    def setUp(self):
        super().setUp()
        self.res = docparse.extract_markdown(self.write("note.md", MD_TEXT))

    def test_headings(self):
        self.assertEqual([(b["text"], b["level"]) for b in self.of_kind(self.res, "heading")],
                         [("近场信道估计", 1), ("实验设置", 2)])

    def test_code_block_not_split_and_not_parsed_as_heading(self):
        codes = self.of_kind(self.res, "code")
        self.assertEqual(len(codes), 1, "代码块必须整块保留，不能被空行切碎")
        self.assertEqual(codes[0]["lang"], "python")
        self.assertIn("# 这一行是代码注释", codes[0]["text"])
        self.assertIn("def loss(x):", codes[0]["text"])
        # 代码里的 # 注释绝不能变成 heading 块
        self.assertNotIn("这一行是代码注释",
                         " ".join(b["text"] for b in self.of_kind(self.res, "heading")))

    def test_table_recognized(self):
        tables = self.of_kind(self.res, "table")
        self.assertEqual(len(tables), 1)
        self.assertEqual(tables[0]["rows"],
                         [["模型", "准确率", "F1"], ["BERT", "0.83", "0.81"],
                          ["Ours", "0.91", "0.90"]], "分隔行要丢掉，表头行要留着")

    def test_frontmatter_into_meta_and_title(self):
        self.assertEqual(self.res["meta"]["frontmatter"]["tags"], "组会")
        self.assertEqual(self.res["title"], "近场信道估计笔记")

    def test_paragraphs_split_on_blank_line(self):
        paras = [b["text"] for b in self.of_kind(self.res, "paragraph")]
        self.assertEqual(paras[0], "正文第一段，讨论球面波模型。\n同一段的第二行。")
        self.assertEqual(paras[-1], "结尾段落。")

    def test_unclosed_fence_is_warned_not_lost(self):
        res = docparse.extract_markdown(self.write("x.md", "# 标题\n\n```\nabc\ndef\n"))
        self.assertEqual(self.of_kind(res, "code")[0]["text"], "abc\ndef")
        self.assertTrue(any("未闭合" in w for w in res["warnings"]))


# ── 纯文本与编码 ──

class TestText(_Base):
    def test_gbk_chinese_file(self):
        """中文环境最常见的线上翻车点：GBK 文本被当成 utf-8 或被 latin-1 静默吃掉。"""
        p = self.write("gbk.txt", TXT_TEXT, encoding="gbk")
        res = docparse.extract_text(p)
        self.assertEqual(res["meta"]["encoding"], "gbk")
        self.assertIn("矩阵、信道、稀疏、重构", res["text"])
        self.assertNotIn("�", res["text"])
        self.assertEqual(res["warnings"], [], "命中 gbk 时不该有兜底 warning")

    def test_utf8_bom_stripped(self):
        p = self.write("bom.txt", TXT_TEXT, encoding="utf-8-sig")
        res = docparse.extract_text(p)
        self.assertEqual(res["meta"]["encoding"], "utf-8-sig")
        self.assertFalse(res["blocks"][0]["text"].startswith("﻿"))

    def test_paragraph_split_and_no_fabricated_title(self):
        p = self.write("a.txt", TXT_TEXT)
        res = docparse.extract_text(p)
        self.assertEqual(len(self.of_kind(res, "paragraph")), 3)
        self.assertIsNone(res["title"], "纯文本没有标题就是没有，不能拿首行冒充")
        self.assertTrue(res["meta"]["first_line"].startswith("第一段中文"))

    def test_latin1_fallback_is_warned(self):
        p = self.tmp / "l1.txt"
        p.write_bytes(b"caf\xe9 na\xefve")     # 既不是 utf-8 也不是合法 gbk 中文
        res = docparse.extract_text(p)
        self.assertEqual(res["meta"]["encoding"], "latin-1")
        self.assertTrue(any("latin-1" in w for w in res["warnings"]))

    def test_binary_like_goes_to_replace_with_warning(self):
        p = self.tmp / "bin.txt"
        p.write_bytes(b"\x00\x01\xff\xfe\x80abc\x00")
        res = docparse.decode_bytes(p.read_bytes())
        self.assertEqual(res[1], "utf-8/replace")
        self.assertTrue(any("编码探测全部失败" in w for w in res[2]))


# ── 格式识别 ──

class TestDetect(_Base):
    def test_extension_wins_when_consistent(self):
        self.assertEqual(docparse.detect_format(self.write("a.md", MD_TEXT)), "markdown")
        self.assertEqual(docparse.detect_format(self.write("a.txt", TXT_TEXT)), "text")
        self.assertEqual(docparse.detect_format(self.write("a.html", HTML_TEXT)), "html")
        self.assertEqual(docparse.detect_format(make_docx(self.tmp / "a.docx")), "word")
        self.assertEqual(docparse.detect_format(make_pptx(self.tmp / "a.pptx")), "powerpoint")

    def test_txt_extension_but_docx_content(self):
        """改扩展名是常态：.txt 里装着 docx（zip 魔数）时以内容为准。"""
        src = Path(make_docx(self.tmp / "real.docx"))
        fake = self.tmp / "disguised.txt"
        fake.write_bytes(src.read_bytes())
        self.assertEqual(docparse.detect_format(fake), "word")
        detail = docparse.detect_format_detail(fake)
        self.assertEqual(detail["by"], "magic")
        self.assertIn("与内容不符", detail["note"])
        # 而且真能按 word 解析出来，不是只把格式名改对
        res = docparse.parse(fake)
        self.assertEqual((res["status"], res["format"]), ("ok", "word"))
        self.assertIn("这是中文正文段落，用于验证编码与顺序。", res["text"])
        self.assertTrue(any("与内容不符" in w for w in res["warnings"]))

    def test_txt_extension_but_pptx_content(self):
        src = Path(make_pptx(self.tmp / "real.pptx"))
        fake = self.tmp / "disguised2.txt"
        fake.write_bytes(src.read_bytes())
        self.assertEqual(docparse.detect_format(fake), "powerpoint")

    def test_docx_extension_but_plain_text_content(self):
        fake = self.write("fake.docx", TXT_TEXT)
        self.assertEqual(docparse.detect_format(fake), "text")
        self.assertIn("不是 ZIP 包", docparse.detect_format_detail(fake)["note"])

    def test_txt_extension_but_html_content(self):
        self.assertEqual(docparse.detect_format(self.write("page.txt", HTML_TEXT)), "html")

    def test_html_extension_without_any_tag(self):
        self.assertEqual(docparse.detect_format(self.write("plain.html", TXT_TEXT)), "text")

    def test_markdown_with_inline_html_stays_markdown(self):
        """Markdown 里内联 HTML 合法，不能因为出现 < 就改判 html。"""
        p = self.write("inline.md", "<span>行内标签</span>\n\n# 标题\n")
        self.assertEqual(docparse.detect_format(p), "markdown")

    def test_pdf_and_xlsx_rejected_with_reason(self):
        pdf = self.tmp / "paper.pdf"
        pdf.write_bytes(b"%PDF-1.7\n1 0 obj\n")
        self.assertIsNone(docparse.detect_format(pdf))
        self.assertIn("pdfimport", docparse.detect_format_detail(pdf)["note"])
        xlsx = self.tmp / "book.xlsx"
        with zipfile.ZipFile(xlsx, "w") as z:
            z.writestr("xl/workbook.xml", "<workbook/>")
        self.assertIsNone(docparse.detect_format(xlsx))
        self.assertIn("Excel", docparse.detect_format_detail(xlsx)["note"])

    def test_unknown_extension_not_guessed(self):
        self.assertIsNone(docparse.detect_format(self.write("data.csv", "a,b\n1,2\n")))
        self.assertFalse(docparse.is_supported(self.write("x.bin", "abc")))


# ── 单文件 parse 的失败语义 ──

class TestParse(_Base):
    def test_missing_file_is_failed_not_exception(self):
        res = docparse.parse(self.tmp / "nope.txt")
        self.assertEqual(res["status"], "failed")
        self.assertEqual(res["reason"], "文件不存在")
        # 失败时键与成功时完全一致，调用方不必写两套取值逻辑
        for k in ("path", "format", "title", "text", "blocks", "meta",
                  "warnings", "n_blocks", "n_chars"):
            self.assertIn(k, res)

    def test_directory_is_failed(self):
        d = self.tmp / "sub"
        d.mkdir()
        self.assertEqual(docparse.parse(d)["status"], "failed")

    def test_empty_file_is_skipped(self):
        p = self.tmp / "empty.txt"
        p.write_bytes(b"")
        res = docparse.parse(p)
        self.assertEqual((res["status"], res["reason"], res["format"]),
                         ("skipped", "空文件", "text"))

    def test_unsupported_is_skipped(self):
        res = docparse.parse(self.write("a.csv", "a,b\n"))
        self.assertEqual(res["status"], "skipped")
        self.assertIn("不支持的格式", res["reason"])

    def test_whitespace_only_is_skipped(self):
        res = docparse.parse(self.write("blank.txt", "\n   \n\t\n"))
        self.assertEqual(res["status"], "skipped")
        self.assertIn("内容为空", res["reason"])

    def test_corrupt_docx_is_failed_with_exception_type(self):
        res = docparse.parse(make_broken_docx(self.tmp / "bad.docx"))
        self.assertEqual((res["status"], res["format"]), ("failed", "word"))
        head = res["reason"].split(":")[0]
        self.assertRegex(head, r"^[A-Za-z_][A-Za-z0-9_.]*$",
                         f"reason 必须以异常类型名开头，实际为 {res['reason']!r}")

    def test_ok_counts(self):
        res = docparse.parse(self.write("note.md", MD_TEXT))
        self.assertEqual(res["status"], "ok")
        self.assertEqual(res["n_blocks"], len(res["blocks"]))
        self.assertEqual(res["n_chars"], len(res["text"]))
        self.assertGreater(res["n_chars"], 0)


# ── 批量三态计数 ──

class TestParseBatch(_Base):
    def build(self) -> list[str]:
        """10 个文件，三态数字可以手算：ok 5 / skipped 3 / failed 2。"""
        note = self.write("note.md", MD_TEXT)
        dup = self.tmp / "note_copy.md"
        dup.write_bytes(note.read_bytes())       # 同内容不同名 → skipped(重复)
        empty = self.tmp / "empty.txt"
        empty.write_bytes(b"")                   # → skipped(空文件)
        return [
            make_docx(self.tmp / "a.docx"),                  # ok  word
            make_pptx(self.tmp / "a.pptx"),                  # ok  powerpoint
            str(note),                                       # ok  markdown
            str(self.write("page.html", HTML_TEXT)),         # ok  html
            str(self.write("gbk.txt", TXT_TEXT, "gbk")),     # ok  text
            str(dup),                                        # skipped markdown
            str(empty),                                      # skipped text
            str(self.write("data.csv", "a,b\n1,2\n")),       # skipped unknown
            make_broken_docx(self.tmp / "bad.docx"),         # failed  word
            str(self.tmp / "ghost.txt"),                     # failed  unknown
        ]

    def test_three_state_counts(self):
        out = docparse.parse_batch(self.build())
        self.assertEqual((out["ok"], out["skipped"], out["failed"]), (5, 3, 2))
        self.assertEqual(len(out["items"]), 10)
        self.assertEqual(out["ok"] + out["skipped"] + out["failed"], len(out["items"]))

    def test_by_format_grouping(self):
        out = docparse.parse_batch(self.build())
        self.assertEqual(out["by_format"], {
            "word": {"ok": 1, "failed": 1, "skipped": 0},
            "powerpoint": {"ok": 1, "failed": 0, "skipped": 0},
            "markdown": {"ok": 1, "failed": 0, "skipped": 1},
            "html": {"ok": 1, "failed": 0, "skipped": 0},
            "text": {"ok": 1, "failed": 0, "skipped": 1},
            "unknown": {"ok": 0, "failed": 1, "skipped": 1},
        })

    def test_one_bad_file_does_not_stop_the_batch(self):
        paths = [str(self.tmp / "ghost.txt"),
                 make_broken_docx(self.tmp / "bad.docx"),
                 str(self.write("note.md", MD_TEXT))]
        out = docparse.parse_batch(paths)
        self.assertEqual([i["status"] for i in out["items"]],
                         ["failed", "failed", "ok"])
        self.assertGreater(out["items"][-1]["n_blocks"], 0)

    def test_duplicate_reason_points_at_the_original(self):
        out = docparse.parse_batch(self.build())
        dup = next(i for i in out["items"] if i["path"].endswith("note_copy.md"))
        self.assertEqual(dup["status"], "skipped")
        self.assertIn("note.md", dup["reason"])
        self.assertIn("重复内容", dup["reason"])

    def test_items_carry_countable_fields(self):
        out = docparse.parse_batch(self.build())
        ok_items = [i for i in out["items"] if i["status"] == "ok"]
        for it in ok_items:
            for k in ("path", "status", "format", "reason", "n_blocks", "n_chars"):
                self.assertIn(k, it)
            self.assertGreater(it["n_blocks"], 0)
            self.assertGreater(it["n_chars"], 0)
        failed = [i for i in out["items"] if i["status"] == "failed"]
        self.assertTrue(all(i["reason"] for i in failed), "failed 必须给出原因")

    def test_empty_list(self):
        out = docparse.parse_batch([])
        self.assertEqual(out, {"ok": 0, "failed": 0, "skipped": 0,
                               "items": [], "by_format": {}})

    def test_progress_callback(self):
        seen = []
        docparse.parse_batch([str(self.write("a.md", MD_TEXT)),
                              str(self.write("b.txt", TXT_TEXT))],
                             progress=lambda f, s, m: seen.append((f, s, m)))
        self.assertEqual(len(seen), 3)                 # 2 个文件 + 收尾
        self.assertEqual(seen[0][0], 0.0)
        self.assertEqual(seen[-1][0], 1.0)
        self.assertEqual(seen[-1][1], "done")
        self.assertIn("成功 2", seen[-1][2])

    def test_broken_progress_callback_does_not_break_batch(self):
        def boom(frac, stage, message):
            raise RuntimeError("回调自己炸了")
        out = docparse.parse_batch([str(self.write("a.md", MD_TEXT))], progress=boom)
        self.assertEqual(out["ok"], 1)

    def test_keep_parsed_opt_in(self):
        paths = [str(self.write("a.md", MD_TEXT))]
        self.assertNotIn("parsed", docparse.parse_batch(paths)["items"][0])
        item = docparse.parse_batch(paths, keep_parsed=True)["items"][0]
        self.assertEqual(item["parsed"]["status"], "ok")
        self.assertTrue(item["parsed"]["blocks"])

    def test_batch_is_deterministic(self):
        """同一批文件连跑两次结果必须完全一致（排序/遍历不能引入漂移）。"""
        paths = self.build()
        a, b = docparse.parse_batch(paths), docparse.parse_batch(paths)
        self.assertEqual([i["path"] for i in a["items"]], [i["path"] for i in b["items"]])
        self.assertEqual([i["status"] for i in a["items"]],
                         [i["status"] for i in b["items"]])
        self.assertEqual(a["by_format"], b["by_format"])


# ── Markdown 中间形态 ──

class TestToMarkdown(_Base):
    def test_docx_to_markdown(self):
        md = docparse.to_markdown(docparse.parse(make_docx(self.tmp / "a.docx")))
        self.assertIn("# 多格式解析测试文档", md)
        self.assertIn("# 绪论", md)
        self.assertIn("## 方法", md)
        self.assertIn("### 小结", md)
        self.assertIn("| 模型 | 准确率 | F1 |", md)
        self.assertIn("| --- | --- | --- |", md)
        self.assertIn("| BERT | 0.83 | 0.81 |", md)

    def test_pptx_to_markdown(self):
        md = docparse.to_markdown(docparse.parse(make_pptx(self.tmp / "a.pptx")))
        self.assertIn("## 第 1 页：研究背景", md)
        self.assertIn("## 第 2 页：实验结果", md)
        self.assertNotIn("第 3 页", md)          # 空白页不产生小节
        self.assertIn("> 备注：", md)
        self.assertIn("| 方法 | NMSE |", md)
        # 标题已经在小节名里，正文不该再重复一遍
        self.assertEqual(md.count("研究背景"), 1)

    def test_markdown_roundtrip_keeps_code_fence(self):
        md = docparse.to_markdown(docparse.parse(self.write("note.md", MD_TEXT)))
        self.assertIn("```python", md)
        self.assertIn("def loss(x):", md)
        self.assertIn("# 这一行是代码注释", md)

    def test_pipe_in_cell_escaped(self):
        parsed = docparse.extract_html("<table><tr><td>a|b</td><td>c</td></tr></table>")
        self.assertIn(r"a\|b", docparse.to_markdown(parsed))

    def test_empty_input(self):
        self.assertEqual(docparse.to_markdown({}), "")
        self.assertEqual(docparse.to_markdown(
            {"blocks": [], "title": None}), "")


# ── 复核阶段补的回归用例 ──
#
# 这一组全部来自「程序化夹具太理想，真实文件不长这样」这条线索：
# python-docx 生成的 docx 永远没有内容控件，python-pptx 的 shapes.title 又刚好和
# 遍历 shapes 拿到的是不同对象，手写的 HTML 片段也总是好心地把 <title> 闭合了。
# 每条都对应一个已在真实形态下复现、并已修掉的缺陷。

class TestRegressions(_Base):

    def test_pptx_title_not_duplicated_in_slide_text(self):
        """python-pptx 每次访问都新建 proxy，`sh is shapes.title` 恒为 False。

        用 `is` 去排除标题占位符，等于没排除——标题会被当普通正文再收一遍。
        注意不能拿 to_markdown 的输出来验：它会把「与小节名相同的行」整行滤掉，
        正好把这个重复掩盖过去（原测试就是这么漏掉的）。必须直接看 slide 块的 text。
        """
        res = docparse.extract_pptx(make_pptx(self.tmp / "a.pptx"))
        s1 = self.of_kind(res, "slide")[0]
        self.assertEqual(s1["text"].count("研究背景"), 1,
                         f"标题被重复收了一遍：{s1['text']!r}")
        self.assertEqual(s1["text"],
                         "研究背景\n近场区域的球面波模型\n远场近似在大孔径下失效")

    def test_unclosed_title_does_not_swallow_the_page(self):
        """`<title>` 忘闭合不能吞掉整篇正文。

        这是 <nav> 那个 bug 的同款：不收口的话整份网页存档会以
        「skipped / 内容为空」的形式无声消失，warnings 里连条线索都没有。
        Python 3.13+ 把 <title> 当 RCDATA，缺 </title> 时解析器连 starttag 回调
        都不会再触发，所以修复点必须在喂进解析器之前。
        """
        p = self.write("clip.html",
                       "<html><head><title>近场信道估计综述\n"
                       "<meta charset='utf-8'></head><body>"
                       "<h1>正文标题</h1><p>整篇正文都在这里。</p>"
                       "<table><tr><td>a</td><td>b</td></tr></table></body></html>")
        res = docparse.parse(p)
        self.assertEqual(res["status"], "ok", f"整篇被吞了：{res['reason']}")
        self.assertIn("整篇正文都在这里。", res["text"])
        self.assertEqual(self.kinds(res), ["heading", "paragraph", "table"])
        self.assertEqual(res["title"], "近场信道估计综述")
        self.assertTrue(any("<title>" in w for w in res["warnings"]),
                        "强制收口这件事必须留痕")

    def test_normal_title_still_wins_and_warns_nothing(self):
        res = docparse.extract_html(
            "<html><head><title>正常标题</title></head><body><p>正文</p></body></html>")
        self.assertEqual(res["title"], "正常标题")
        self.assertEqual(res["warnings"], [])
        self.assertEqual([b["text"] for b in res["blocks"]], ["正文"])

    def test_docx_content_control_body_is_not_lost(self):
        """Word 的自动目录 / 封面 / 模板占位区把正文包在 `<w:sdt><w:sdtContent>` 里。

        python-docx **生成**的文件永远不含 sdt，所以只用程序化夹具是测不出来的；
        真实投稿稿里丢的是整章内容，而且一声不吭。顺带验证下钻后顺序仍是文档顺序。
        """
        doc = Document()
        doc.add_heading("绪论", level=1)
        last = doc.add_paragraph("最后一段")
        sdt = parse_xml(
            f"<w:sdt {nsdecls('w')}><w:sdtContent><w:p><w:r>"
            f"<w:t>内容控件里的正文</w:t></w:r></w:p></w:sdtContent></w:sdt>")
        last._p.addprevious(sdt)
        doc.save(str(self.tmp / "sdt.docx"))
        res = docparse.extract_docx(self.tmp / "sdt.docx")
        self.assertEqual([b["text"] for b in res["blocks"]],
                         ["绪论", "内容控件里的正文", "最后一段"])

    def test_nul_bytes_are_reported_and_stripped(self):
        """NUL 是合法 UTF-8 码位，探测链不会失败——二进制改名成 .txt 会静默通过。

        原实现里这些字节会一路流进 blocks，最后进 sqlite / FTS。
        """
        p = self.tmp / "sneaky.txt"
        p.write_bytes("第一段中文".encode() + b"\x00\x07" + "还有下文".encode())
        res = docparse.parse(p)
        self.assertEqual(res["status"], "ok")
        self.assertNotIn("\x00", res["text"])
        self.assertNotIn("\x07", res["text"])
        self.assertIn("第一段中文还有下文", res["text"])
        self.assertTrue(any("控制字符" in w for w in res["warnings"]),
                        "剔除了就要说，不能静默改写正文")

    def test_control_chars_stripped_in_code_blocks_too(self):
        """code / pre 走的是不经过 _collapse 的路径，剔除点必须放在更下游。"""
        res = docparse.extract_markdown(
            self.write("c.md", "```\nprint(1)\x00\n```\n"))
        self.assertNotIn("\x00", self.of_kind(res, "code")[0]["text"])

    def test_oversized_file_is_not_hashed_before_the_size_gate(self):
        """去重哈希跑在体积闸门之前：不设上限的话，超限文件会被完整读一遍才丢掉。"""
        p = self.write("big.txt", "x" * 5000)
        with mock.patch.object(docparse, "MAX_FILE_BYTES", 1000):
            self.assertIsNone(docparse._sha256_file(p))
            out = docparse.parse_batch([str(p)])
        self.assertEqual(out["skipped"], 1)
        self.assertIn("超过", out["items"][0]["reason"])

    def test_to_markdown_round_trips_through_extract_markdown(self):
        """to_markdown 是本模块对下游承诺的「统一中间形态」，它必须能被自己读回来。

        单元格里含 `|` 时 to_markdown 会写成 `\\|`，而 _md_row 原来无脑 split("|")，
        两列的表会被读成三列——下游按列取值全错位。
        """
        parsed = docparse.extract_html(
            "<table><tr><th>模型</th><th>说明</th></tr>"
            "<tr><td>Ours</td><td>a|b 混合</td></tr></table>")
        back = docparse.extract_markdown(docparse.to_markdown(parsed))
        self.assertEqual(self.of_kind(back, "table")[0]["rows"],
                         [["模型", "说明"], ["Ours", "a|b 混合"]])

    def test_markdown_inline_text_whose_last_line_looks_like_a_filename(self):
        """`Path("# 标题\\n\\n详见 report.md\\n").suffix` == ".md\\n"。

        原来的「多行 + 无后缀 = 直接喂的文本」判据会把这段正文当路径去 open()，抛 OSError。
        """
        res = docparse.extract_markdown("# 标题\n\n详见 report.md 里的说明\n")
        self.assertEqual([b["text"] for b in res["blocks"]],
                         ["标题", "详见 report.md 里的说明"])

    def test_slide_body_line_equal_to_title_is_not_deleted(self):
        """去重只该掐掉首行的标题回声，不该把正文里合法重复的一行也删掉。"""
        parsed = {"title": None, "blocks": [
            {"kind": "slide", "index": 1, "title": "结论", "has_title": True,
             "text": "结论\n中间正文\n结论"}]}
        md = docparse.to_markdown(parsed)
        self.assertIn("中间正文", md)
        self.assertEqual(md.count("结论"), 2, f"正文里那一行被误删了：{md!r}")

    def test_utf16_bom_file_round_trips(self):
        """带 BOM 的 UTF-16 要走单独分支，不能被 latin-1 兜成夹 \\x00 的乱码。"""
        p = self.tmp / "u16.txt"
        p.write_bytes(TXT_TEXT.encode("utf-16"))
        res = docparse.extract_text(p)
        self.assertEqual(res["meta"]["encoding"], "utf-16")
        self.assertIn("矩阵、信道、稀疏、重构", res["text"])
        self.assertEqual(res["warnings"], [])


if __name__ == "__main__":
    unittest.main()
