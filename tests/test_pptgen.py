import shutil
import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from papernest import config, db, pptgen


def _seed_paper(c, norm_key, title, card):
    pid = db.insert_l0(c, {"norm_key": norm_key, "title": title,
                           "abstract": "This paper studies near-field channel estimation for XL-MIMO. "
                                       "We propose a deep generative model. Simulations show gains.",
                           "year": 2025, "venue": "IEEE TWC", "authors": ["A. Author", "B. Author"],
                           "doi": None, "arxiv_id": "2501.00001", "source": "s2"})
    db.save_card(c, pid, card, "test-model")
    c.commit()
    return pid


CARD = {
    "tldr": "提出近场生成式信道估计模型",
    "problem": "近场球面波前让远场估计方法失效",
    "method": "深度生成模型 + 可见区域先验",
    "results": "MSE 相对基线下降 30%",
    "key_findings": [
        {"claim": "MSE 相对最小二乘基线下降约 30%", "page": 5, "verified": True},
        {"claim": "复杂度随天线数线性增长", "page": 7, "verified": False},
    ],
    "limitations": "仅在仿真中验证",
    "relation_to_topic": "可直接用于波束训练评估基线",
}


class PptGenTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.mkdtemp(prefix="papernest_ppt_")
        self.addCleanup(shutil.rmtree, self._tmp, True)
        self._old_db, self._old_data = config.DB_PATH, config.DATA_DIR
        config.DB_PATH = Path(self._tmp) / "test.db"
        db.init_db()
        with db.conn() as c:
            self.pid = _seed_paper(c, "arxiv:2501.00001", "Deep Generative Near-Field Estimation", CARD)

    def tearDown(self):
        config.DB_PATH, config.DATA_DIR = self._old_db, self._old_data

    def _slides_text(self, path):
        prs = Presentation(str(path))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text_frame.text)
        return "\n".join(texts)

    def test_single_paper_deck_structure(self):
        out = Path(self._tmp) / "out"
        r = pptgen.deck([self.pid], topic="XL-MIMO 近场", out_dir=out)
        self.assertTrue(Path(r["path"]).exists())
        self.assertGreaterEqual(r["slides"], 5)  # 封面+背景+方法+结果+局限+致谢
        text = self._slides_text(r["path"])
        for expected in ("Deep Generative Near-Field Estimation", "研究背景与问题",
                         "关键结果", "✓", "？", "p5"):
            self.assertIn(expected, text, f"幻灯片缺少：{expected}")

    def test_multi_paper_deck_has_overview_and_refs(self):
        with db.conn() as c:
            pid2 = _seed_paper(c, "arxiv:2501.00002", "Second Paper on Beam Training", CARD)
        out = Path(self._tmp) / "out2"
        # narrative=False：测试不调真模型
        r = pptgen.deck([self.pid, pid2], topic="", out_dir=out, narrative=False)
        text = self._slides_text(r["path"])
        self.assertIn("文献总览", text)
        self.assertIn("参考文献", text)
        self.assertIn("Second Paper on Beam Training", text)

    def test_empty_ids_rejected(self):
        with self.assertRaises(ValueError):
            pptgen.deck([], out_dir=Path(self._tmp))


if __name__ == "__main__":
    unittest.main()
