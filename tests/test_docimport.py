"""多格式文档入库：docparse 解析结果 → 论文库（去重、pages/chunks、三态计数）。

夹具全部用 python-docx / python-pptx **真的生成文件**，不是伪造 dict——
本项目踩过两次「合成夹具掩盖真实数据缺陷」的亏，解析类模块的测试必须过真实文件。
"""
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from papernest import config, db, docimport

LONG = "检索增强生成把外部知识引入生成过程，评测要同时量召回与忠实度。" * 12


def make_docx(d: Path) -> Path:
    import docx
    doc = docx.Document()
    doc.add_heading("检索增强生成的评测方法", 0)
    doc.add_heading("1 引言", level=1)
    doc.add_paragraph(LONG)
    doc.add_heading("2 方法", level=1)
    doc.add_paragraph("我们提出两阶段评测框架。" * 20)
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "方法"
    t.cell(0, 1).text = "Recall"
    t.cell(1, 0).text = "混合检索"
    t.cell(1, 1).text = "0.771"
    p = d / "a.docx"
    doc.save(p)
    return p


def make_pptx(d: Path) -> Path:
    import pptx
    pr = pptx.Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[1])
    s.shapes.title.text = "组会汇报"
    s.placeholders[1].text = LONG
    s.notes_slide.notes_text_frame.text = "备注：要强调那个负结果。" * 10
    s2 = pr.slides.add_slide(pr.slide_layouts[1])
    s2.shapes.title.text = "下一步"
    s2.placeholders[1].text = "接入 Cross-Encoder 做 A/B。" * 15
    p = d / "b.pptx"
    pr.save(p)
    return p


class DocImportTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.mkdtemp(prefix="papernest_di_")
        self.addCleanup(shutil.rmtree, self._tmp, True)
        self.d = Path(self._tmp)
        self._old_db, self._old_dir = config.DB_PATH, config.DATA_DIR
        config.DATA_DIR = self.d / "data"
        config.DB_PATH = config.DATA_DIR / "test.db"
        self._patches = [
            mock.patch("papernest.embeddings.available", return_value=False),
            mock.patch("papernest.llm.available", return_value=False),
        ]
        for p in self._patches:
            p.start()
        db.init_db()

    def tearDown(self):
        for p in self._patches:
            p.stop()
        config.DB_PATH, config.DATA_DIR = self._old_db, self._old_dir

    # ── 单份入库 ──

    def test_docx_import_populates_pages_and_chunks(self):
        r = docimport.import_document(make_docx(self.d), make_card=False)
        self.assertEqual(r["status"], "ok")
        self.assertEqual(r["format"], "word")
        self.assertGreater(r["chunks"], 0)
        with db.conn() as c:
            row = c.execute("SELECT level, source FROM papers WHERE id=?",
                            (r["paper_id"],)).fetchone()
            self.assertEqual(row["level"], 2)          # 有全文就是 L2
            self.assertEqual(row["source"], "upload-word")

    def test_heading_becomes_section_path(self):
        """检索命中一个块时要知道它属于哪一节——与 PDF 的章节块对齐。"""
        r = docimport.import_document(make_docx(self.d), make_card=False)
        with db.conn() as c:
            paths = [x["section_path"] for x in c.execute(
                "SELECT section_path FROM chunks WHERE paper_id=?", (r["paper_id"],))]
        self.assertIn("2 方法", paths)

    def test_docx_table_becomes_table_chunk(self):
        r = docimport.import_document(make_docx(self.d), make_card=False)
        with db.conn() as c:
            n = c.execute("SELECT COUNT(1) n FROM chunks WHERE paper_id=? AND kind='table'",
                          (r["paper_id"],)).fetchone()["n"]
        self.assertGreaterEqual(n, 1)

    def test_pptx_slides_become_pages_and_notes_are_kept(self):
        r = docimport.import_document(make_pptx(self.d), make_card=False)
        self.assertEqual(r["status"], "ok")
        self.assertEqual(r["pages_stored"], 2)          # 一张幻灯片算一页
        with db.conn() as c:
            hit = c.execute("SELECT 1 FROM pages WHERE paper_id=? AND text LIKE '%负结果%'",
                            (r["paper_id"],)).fetchone()
        self.assertIsNotNone(hit, "演讲者备注里常有真正的解释，不能丢")

    def test_markdown_and_html_and_gbk_text(self):
        (self.d / "n.md").write_text("# 笔记标题\n\n" + LONG, encoding="utf-8")
        (self.d / "p.html").write_text(
            "<html><head><title>网页标题</title></head><body>"
            "<script>var x=1;</script><h1>正文</h1><p>" + LONG + "</p></body></html>",
            encoding="utf-8")
        (self.d / "g.txt").write_text(LONG, encoding="gbk")   # 中文环境最常见的坑
        for name, fmt in (("n.md", "markdown"), ("p.html", "html"), ("g.txt", "text")):
            r = docimport.import_document(self.d / name, make_card=False)
            self.assertEqual(r["status"], "ok", f"{name}: {r.get('reason')}")
            self.assertEqual(r["format"], fmt)
        with db.conn() as c:
            n = c.execute("SELECT COUNT(1) n FROM chunks WHERE text LIKE '%var x=1%'"
                          ).fetchone()["n"]
        self.assertEqual(n, 0, "script 内容不能进检索索引")

    def test_body_only_content_is_searchable(self):
        docimport.import_document(make_docx(self.d), make_card=False)
        with db.conn() as c:
            self.assertTrue(db.search_chunks_fts(c, "两阶段评测框架", 3))

    # ── 去重与三态 ──

    def test_same_document_twice_is_skipped(self):
        p = make_docx(self.d)
        first = docimport.import_document(p, make_card=False)
        second = docimport.import_document(p, make_card=False)
        self.assertEqual(second["status"], "skipped")
        self.assertTrue(second["duplicate"])
        self.assertEqual(second["paper_id"], first["paper_id"])

    def test_too_short_document_is_skipped_not_failed(self):
        (self.d / "tiny.txt").write_text("太短了", encoding="utf-8")
        r = docimport.import_document(self.d / "tiny.txt", make_card=False)
        self.assertEqual(r["status"], "skipped")
        self.assertIn("入库下限", r["reason"])

    def test_missing_file_is_failed_with_reason(self):
        r = docimport.import_document(self.d / "nope.docx", make_card=False)
        self.assertEqual(r["status"], "failed")
        self.assertTrue(r["reason"])

    def test_corrupt_docx_is_failed_not_crash(self):
        (self.d / "bad.docx").write_bytes(b"PK\x03\x04not really a docx")
        r = docimport.import_document(self.d / "bad.docx", make_card=False)
        self.assertEqual(r["status"], "failed")

    # ── 批量 ──

    def test_batch_counts_are_three_state_and_add_up(self):
        make_docx(self.d)
        make_pptx(self.d)
        (self.d / "tiny.txt").write_text("短", encoding="utf-8")
        paths = [self.d / "a.docx", self.d / "b.pptx", self.d / "tiny.txt",
                 self.d / "missing.md"]
        r = docimport.import_batch(paths, make_card=False)
        self.assertEqual(r["total"], 4)
        self.assertEqual(r["ok"] + r["skipped"] + r["failed"], 4)
        self.assertEqual(r["ok"], 2)
        self.assertEqual(r["skipped"], 1)     # tiny
        self.assertEqual(r["failed"], 1)      # missing
        self.assertEqual(r["by_format"]["word"]["ok"], 1)

    def test_one_bad_file_does_not_stop_the_batch(self):
        make_docx(self.d)
        (self.d / "bad.docx").write_bytes(b"PK\x03\x04garbage")
        r = docimport.import_batch([self.d / "bad.docx", self.d / "a.docx"],
                                   make_card=False)
        self.assertEqual(r["ok"], 1)
        self.assertEqual(r["failed"], 1)

    def test_collect_paths_skips_pdf(self):
        make_docx(self.d)
        (self.d / "x.pdf").write_bytes(b"%PDF-1.4 stub")
        got = {p.name for p in docimport.collect_paths([self.d])}
        self.assertIn("a.docx", got)
        self.assertNotIn("x.pdf", got, "PDF 走 import-pdf，不该被这条路径收走")

    def test_no_card_means_no_llm_call(self):
        """--no-cards 必须是真的 0 token。"""
        with mock.patch("papernest.cards.make_card") as m:
            docimport.import_document(make_docx(self.d), make_card=False)
        m.assert_not_called()


if __name__ == "__main__":
    unittest.main()
